"""Bounded local parsing; bank spreadsheets never pass through generic uploads/OCR."""
import base64
import csv
import io
from pathlib import PurePosixPath
import re
import zipfile

from .errors import ContractError


def masked_cell(value):
    return re.sub(r'\b[A-Z]{2}[0-9]{2}(?:[ ]?[A-Z0-9]){10,30}\b','[dato bancario protegido]',str(value),flags=re.I)


def parse_workbook(encoded, filename, sheet_index=0):
    if not isinstance(encoded,str) or len(encoded)>24*1024*1024:
        raise ContractError('El archivo supera el limite de 16 MB.')
    try:raw=base64.b64decode(encoded,validate=True)
    except Exception:raise ContractError('Archivo no valido.') from None
    if not raw or len(raw)>16*1024*1024:raise ContractError('Archivo vacio o demasiado grande.')
    if type(sheet_index) is not int or sheet_index<0:raise ContractError('Selecciona una hoja valida.')
    extension=PurePosixPath(str(filename).replace('\\','/')).suffix.lower()
    sheets=[];rows=[]
    try:
        if extension=='.csv':
            content=raw.decode('utf-8-sig')
            try:dialect=csv.Sniffer().sniff(content[:8192],delimiters=',;\t')
            except csv.Error:dialect=csv.excel
            sheets=['Datos'];iterator=csv.reader(io.StringIO(content),dialect)
            for index,row in enumerate(iterator):
                if index>=1016 or len(row)>100:raise ContractError('Maximo 1.000 filas de datos y 100 columnas.')
                rows.append(row)
        elif extension=='.xlsx':
            from defusedxml import ElementTree
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                files=archive.infolist()
                if len(files)>1000 or sum(f.file_size for f in files)>32*1024*1024:
                    raise ContractError('Libro comprimido excesivo.')
                names=set()
                for info in files:
                    file=PurePosixPath(info.filename)
                    if info.filename in names or file.is_absolute() or '..' in file.parts or info.flag_bits&1 or info.file_size>8*1024*1024:
                        raise ContractError('Estructura del libro no admitida.')
                    names.add(info.filename)
                    if 'externalLinks/' in info.filename or info.filename.endswith('vbaProject.bin'):
                        raise ContractError('No se admiten enlaces externos ni macros.')
                    if info.filename.endswith(('.xml','.rels')):
                        # Parse every XML part with entity/DTD protection before the workbook library.
                        ElementTree.fromstring(archive.read(info),forbid_dtd=True,forbid_entities=True,forbid_external=True)
            import openpyxl
            workbook=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
            try:
                sheets=workbook.sheetnames
                if sheet_index>=len(sheets):raise ContractError('La hoja no existe.')
                sheet=workbook.worksheets[sheet_index]
                if (sheet.max_row or 0)>1016 or (sheet.max_column or 0)>100:
                    raise ContractError('Maximo 1.000 filas de datos y 100 columnas.')
                for row in sheet.iter_rows(max_row=1016,max_col=min(sheet.max_column or 100,100)):
                    values=[]
                    for cell in row:
                        if cell.data_type=='f':values.append('[Formula: sustituir por valor]')
                        else:values.append('' if cell.value is None else str(cell.value))
                    rows.append(values)
            finally:workbook.close()
        else:raise ContractError('Utiliza Excel .xlsx o CSV UTF-8. Conserva el archivo original.')
        header_index=next((i for i,row in enumerate(rows[:15]) if any(str(x).strip() for x in row)),None)
        if header_index is None:raise ContractError('No se han encontrado cabeceras.')
        if any(len(str(v))>10000 for row in rows for v in row) or any(len(str(v))>250 for v in rows[header_index]):
            raise ContractError('Hay celdas o cabeceras excesivamente largas.')
        headers=[str(v).strip() or 'Columna '+str(i+1) for i,v in enumerate(rows[header_index])]
        values=[{'row':i+1,'values':[str(v) for v in row]} for i,row in enumerate(rows) if i>header_index and any(str(x).strip() for x in row)]
        if not values or len(values)>1000:raise ContractError('El libro debe contener entre 1 y 1.000 filas de datos.')
        return {'sheets':sheets,'sheet_index':sheet_index,'headers':headers,'rows':values}
    except ContractError:raise
    except Exception:raise ContractError('No se pudo interpretar el libro de forma segura. Revisa el formato y las cabeceras.') from None
