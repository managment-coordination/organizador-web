"""Transfer, evidence corrections and legacy staging through the ERP 5 boundary."""

from .contracts import canonical_json
from .errors import ContractError, ConflictError, NotFoundError
from .receivables_contracts import cents,day,identity,require_fields,text
from .receivables_service import ReceivablesService


class ReconciliationOperations:
    def statement_original(self,session,community,import_id,reason):
        import base64
        from contextlib import closing
        from pathlib import PurePosixPath
        import uuid
        from .audit import write_event
        from .database import connect,write_transaction
        from .migrations import utc_now
        from .banking_service import BankingService
        self._recent_auth(session);reason=text(reason,'Motivo del acceso',maximum=1000)
        with closing(connect(self.database_path)) as conn,write_transaction(conn):
            actor,current=self._session(conn,session,community,'read')
            BankingService._session(conn,session,community,'export');BankingService._session(conn,session,community,'reveal')
            from access_control import permission
            if not permission(current,community,'puede_ver_documentos'):raise PermissionError('No tienes acceso a documentos.')
            row=self._row(conn,'erp_extractos_importaciones',community,import_id)
            data=self.vault.get(conn,community,row['secret_id'],'bank-statement')['original']
            raw=base64.b64decode(data['content_base64'],validate=True) if data.get('content_base64') else canonical_json(data.get('rows',[])).encode()
            extension=PurePosixPath(str(data.get('filename','manual.json'))).suffix.lower().lstrip('.')
            if extension not in ('csv','xls','xlsx','txt','xml','json'):extension='txt'
            now=utc_now();context=self.vault.put(conn,community,'bank-original-access',{'reason':reason,'import_id':row['id']},now)
            write_event(conn,community_id=community,actor=actor,action='erp5.statement.original',entity_type='erp5_import',entity_id=row['id'],
                before=None,after={'import_id':row['id']},reason='Descarga bancaria autorizada',origin='web',request_id=uuid.uuid4().hex,
                entity_version=row['version'],metadata={'protected_context_id':context})
            return {'ok':True,'content_base64':base64.b64encode(raw).decode(),'extension':extension}

    def _transfer_plan(self,conn,community,p):
        require_fields(p,('source_id','destination_id','amount_cents','legs'),('description',))
        source=self._account(conn,community,p['source_id']);destination=self._account(conn,community,p['destination_id'])
        if source['id']==destination['id']:raise ContractError('Selecciona dos cuentas distintas.')
        n=cents(p['amount_cents'],positive=True);legs=[];totals={'source':0,'destination':0}
        if not isinstance(p['legs'],list) or not 1<=len(p['legs'])<=200:raise ContractError('Incluye extremos acreditados de la transferencia.')
        seen=set()
        for leg in p['legs']:
            require_fields(leg,('side','amount_cents'),('movement_id','effective_on','cash_evidence'))
            side=leg['side']
            if side not in totals:raise ContractError('Extremo de transferencia no valido.')
            account=source if side=='source' else destination;amount=cents(leg['amount_cents'],positive=True)
            totals[side]+=amount
            if totals[side]>n:raise ContractError('Los extremos superan el principal.')
            mid=leg.get('movement_id');version=None;remaining=None
            if account['kind']=='banco':
                m=self._row(conn,'erp_banco_movimientos',community,mid)
                if m['id'] in seen:raise ContractError('Movimiento repetido en la transferencia.')
                seen.add(m['id']);version=m['version'];effective=m['operation_on']
                if m['treasury_id']!=account['id'] or m['source_state']!='booked':raise ContractError('Movimiento de otra cuenta o no asentado.')
                if conn.execute('SELECT 1 FROM erp_banco_movimiento_correcciones WHERE id_comunidad=? AND movement_id=?',(community,m['id'])).fetchone():
                    raise ConflictError('Movimiento rectificado; utiliza su sustituto.')
                self._guard_open(conn,community,account['id'],effective)
                remain=self._remaining(conn,community,m)
                remaining=str(remain)
                if (side=='source' and remain>=0) or (side=='destination' and remain<=0) or amount>abs(remain):
                    raise ConflictError('Importe o signo de transferencia incompatible con el remanente.')
            else:
                if mid:raise ContractError('Una evidencia de caja no es movimiento bancario.')
                effective=day(leg['effective_on']);text(leg.get('cash_evidence'),'Justificante de caja')
            ReceivablesService._date_open(conn,community,effective)
            legs.append({'side':side,'treasury_id':account['id'],'movement_id':mid,'version':version,'remaining_cents':remaining,
                         'amount_cents':str(-amount if side=='source' else amount),'effective_on':effective})
        plan={'source_id':source['id'],'destination_id':destination['id'],'amount_cents':str(n),'legs':legs,
              'in_transit_cents':str(n-min(totals.values())),'state':'confirmada' if min(totals.values())==n else 'en_transito'}
        plan['preview_hash']=self.vault.fingerprint(community,'transfer-preview',{'payload':p,'plan':plan})
        return plan

    def transfer_preview(self,session,env):
        return self._write(session,env,'manage_transfers',lambda conn,actor,e,now:self._proposal(conn,actor,e,now,'transfer',e.payload,self._transfer_plan(conn,e.community_id,e.payload)))

    def transfer_confirm(self,session,env):
        def op(conn,actor,e,now):
            proposal=self._row(conn,'erp_conciliacion_propuestas',e.community_id,e.payload['proposal_id']);self._version(proposal,e.expected_version)
            if proposal['state']!='pendiente' or proposal['kind']!='transfer':raise ConflictError('Propuesta de transferencia no disponible.')
            data=self.vault.get(conn,e.community_id,proposal['secret_id'],'reconciliation-proposal');plan=self._transfer_plan(conn,e.community_id,data['payload'])
            if plan['preview_hash']!=data['public']['preview_hash']:raise ConflictError('Han cambiado los extremos de la transferencia.')
            secret=self.vault.put(conn,e.community_id,'treasury-transfer',{'plan':plan,'payload':data['payload']},now)
            tid=conn.execute('INSERT INTO erp_tesoreria_transferencias (id_comunidad,source_id,destination_id,amount_cents,secret_id,registered_at,actor_id) VALUES (?,?,?,?,?,?,?)',
                             (e.community_id,plan['source_id'],plan['destination_id'],int(plan['amount_cents']),secret,now,actor.user_id)).lastrowid
            matchsecret=self.vault.put(conn,e.community_id,'bank-match',{'transfer_id':tid,'plan':plan},now)
            match=conn.execute('INSERT INTO erp_conciliaciones (id_comunidad,proposal_id,secret_id,registered_at,actor_id) VALUES (?,?,?,?,?)',
                               (e.community_id,proposal['id'],matchsecret,now,actor.user_id)).lastrowid
            self._save_transfer_legs(conn,actor,e,now,tid,match,plan['legs'])
            self._event(conn,e,'transfer.confirmed','transfer',tid,{'effect':'link','state':plan['state'],'source_id':plan['source_id'],'destination_id':plan['destination_id']})
            conn.execute("UPDATE erp_conciliacion_propuestas SET state='confirmada',version=version+1 WHERE id=?",(proposal['id'],))
            return {'id':tid,'match_id':match,'state':plan['state'],'in_transit_cents':plan['in_transit_cents']}
        return self._write(session,env,'manage_transfers',op)

    def _save_transfer_legs(self,conn,actor,e,now,tid,match,legs):
        for leg in legs:
            secret=self.vault.put(conn,e.community_id,'transfer-leg',leg,now)
            lid=conn.execute('''INSERT INTO erp_tesoreria_transferencia_extremos
                (id_comunidad,transfer_id,match_id,side,treasury_id,movement_id,amount_cents,effective_on,secret_id,registered_at,actor_id)
                VALUES (?,?,?,?,?,?,?,?,?,?,?)''',(e.community_id,tid,match,leg['side'],leg['treasury_id'],leg['movement_id'],
                int(leg['amount_cents']),leg['effective_on'],secret,now,actor.user_id)).lastrowid
            if leg['movement_id']:
                conn.execute('''INSERT INTO erp_conciliacion_componentes (id_comunidad,match_id,movement_id,kind,transfer_id,amount_cents,registered_at,actor_id)
                    VALUES (?,?,?,'transfer',?,?,?,?)''',(e.community_id,match,leg['movement_id'],tid,int(leg['amount_cents']),now,actor.user_id))
            self._event(conn,e,'transfer.leg_confirmed','transfer-leg',lid,
                {'effect':'monetary','economic_fact_id':'erp5:transfer:'+str(tid),'component_id':str(lid),
                 'treasury_id':leg['treasury_id'],'amount_cents':leg['amount_cents'],'currency':'EUR','effective_on':leg['effective_on']})

    def _transfer_completion(self,conn,community,p):
        require_fields(p,('transfer_id','legs'))
        transfer=self._row(conn,'erp_tesoreria_transferencias',community,p['transfer_id'])
        existing=[dict(r) for r in conn.execute('SELECT * FROM erp_tesoreria_transferencia_extremos WHERE id_comunidad=? AND transfer_id=? ORDER BY id',
                                              (community,transfer['id']))]
        plan=self._transfer_plan(conn,community,{'source_id':transfer['source_id'],'destination_id':transfer['destination_id'],
            'amount_cents':str(transfer['amount_cents']),'legs':p['legs']})
        totals={side:sum(abs(r['amount_cents']) for r in existing if r['side']==side)+sum(abs(int(r['amount_cents'])) for r in plan['legs'] if r['side']==side)
                for side in ('source','destination')}
        if max(totals.values())>transfer['amount_cents']:raise ConflictError('Ese extremo ya esta acreditado o supera el principal.')
        plan.update(transfer_id=transfer['id'],transfer_version=len(existing),state='confirmada' if min(totals.values())==transfer['amount_cents'] else 'en_transito',
                    in_transit_cents=str(transfer['amount_cents']-min(totals.values())))
        plan['preview_hash']=self.vault.fingerprint(community,'transfer-completion',{'plan':plan,'existing':existing})
        return plan

    def transfer_complete_preview(self,session,env):
        return self._write(session,env,'manage_transfers',lambda conn,actor,e,now:self._proposal(conn,actor,e,now,'transfer_completion',e.payload,
            self._transfer_completion(conn,e.community_id,e.payload)))

    def transfer_complete_confirm(self,session,env):
        def op(conn,actor,e,now):
            proposal=self._row(conn,'erp_conciliacion_propuestas',e.community_id,e.payload['proposal_id']);self._version(proposal,e.expected_version)
            if proposal['state']!='pendiente' or proposal['kind']!='transfer_completion':raise ConflictError('Propuesta no disponible.')
            data=self.vault.get(conn,e.community_id,proposal['secret_id'],'reconciliation-proposal')
            plan=self._transfer_completion(conn,e.community_id,data['payload'])
            if plan['preview_hash']!=data['public']['preview_hash']:raise ConflictError('Han cambiado los extremos; revisa de nuevo.')
            secret=self.vault.put(conn,e.community_id,'bank-match',plan,now)
            match=conn.execute('INSERT INTO erp_conciliaciones (id_comunidad,proposal_id,secret_id,registered_at,actor_id) VALUES (?,?,?,?,?)',
                (e.community_id,proposal['id'],secret,now,actor.user_id)).lastrowid
            self._save_transfer_legs(conn,actor,e,now,plan['transfer_id'],match,plan['legs'])
            conn.execute("UPDATE erp_conciliacion_propuestas SET state='confirmada',version=version+1 WHERE id=?",(proposal['id'],))
            self._event(conn,e,'transfer.completed','transfer-completion',match,{'effect':'link','transfer_id':plan['transfer_id'],'state':plan['state']})
            return {'id':plan['transfer_id'],'match_id':match,'state':plan['state'],'in_transit_cents':plan['in_transit_cents']}
        return self._write(session,env,'manage_transfers',op)

    def transfer_list(self,session,q):
        def op(conn,q):
            items=[]
            for r in conn.execute('SELECT * FROM erp_tesoreria_transferencias WHERE id_comunidad=? ORDER BY id DESC',(q.community_id,)):
                totals={side:sum(abs(x[0]) for x in conn.execute('SELECT amount_cents FROM erp_tesoreria_transferencia_extremos WHERE id_comunidad=? AND transfer_id=? AND side=?',
                    (q.community_id,r['id'],side))) for side in ('source','destination')}
                transit=r['amount_cents']-min(totals.values())
                items.append({'id':r['id'],'source_id':r['source_id'],'destination_id':r['destination_id'],
                              'amount_cents':str(r['amount_cents']),'state':'en_transito' if transit else 'confirmada','in_transit_cents':str(transit)})
            return {'items':items}
        return self._read(session,q,'read',op)

    def movement_correct(self,session,env):
        def op(conn,actor,e,now):
            p=require_fields(e.payload,('movement_id','replacement_id','kind'))
            m=self._row(conn,'erp_banco_movimientos',e.community_id,p['movement_id']);self._version(m,e.expected_version)
            replacement=self._row(conn,'erp_banco_movimientos',e.community_id,p['replacement_id'])
            if m['id']==replacement['id'] or m['treasury_id']!=replacement['treasury_id']:raise ContractError('Sustituto no valido.')
            self._guard_open(conn,e.community_id,m['treasury_id'],m['operation_on'])
            if self._remaining(conn,e.community_id,m)!=m['amount_cents']:raise ConflictError('Desconcilia o rectifica los enlaces primero.')
            if p['kind'] not in ('duplicate','substitute'):raise ContractError('Tipo de correccion no valido.')
            if conn.execute('SELECT 1 FROM erp_banco_movimiento_correcciones WHERE id_comunidad=? AND movement_id IN (?,?)',
                            (e.community_id,m['id'],replacement['id'])).fetchone():raise ConflictError('Movimiento ya rectificado.')
            secret=self.vault.put(conn,e.community_id,'bank-correction',{'reason':e.reason,'payload':p},now)
            rid=conn.execute('INSERT INTO erp_banco_movimiento_correcciones (id_comunidad,movement_id,replacement_id,kind,secret_id,registered_at,actor_id) VALUES (?,?,?,?,?,?,?)',
                             (e.community_id,m['id'],replacement['id'],p['kind'],secret,now,actor.user_id)).lastrowid
            return {'id':rid,'movement_id':m['id'],'replacement_id':replacement['id'],'effect':'evidence'}
        return self._write(session,env,'correct',op)

    def legacy_preview(self,session,env):
        def op(conn,actor,e,now):
            self._session(conn,session,e.community_id,'import')
            p=require_fields(e.payload,('treasury_id','rows'))
            self._account(conn,e.community_id,p['treasury_id'])
            if not isinstance(p['rows'],list) or not 1<=len(p['rows'])<=1000:raise ContractError('Revisa hasta 1.000 correspondencias.')
            rows=[]
            # Unattributed legacy rows remain pending; a user-selected target is not ownership proof.
            allowed={'cf_extractos_banco_lineas':('id_linea_banco',)}
            for r in p['rows']:
                require_fields(r,('table','id','movement'))
                if r['table'] not in allowed:raise ContractError('Fuente historica no soportada.')
                key=allowed[r['table']][0]
                columns={x[1] for x in conn.execute(f'PRAGMA table_info({r["table"]})')}
                if 'id_comunidad' not in columns:raise ConflictError('La fuente historica no acredita comunidad. Resuelve primero su procedencia.')
                original=conn.execute(f'SELECT * FROM {r["table"]} WHERE {key}=? AND id_comunidad=?',(identity(r['id']),e.community_id)).fetchone()
                if not original:raise NotFoundError('Fila historica no disponible.')
                from .reconciliation_adapters import normalize,amount
                normalized=normalize(r['movement'])
                if normalized['amount_cents']!=amount(str(original['importe'])):
                    raise ConflictError('El importe historico no coincide exactamente. Conserva la incidencia; no se redondea ni reinterpreta.')
                rows.append({'table':r['table'],'id':str(r['id']),'movement':normalized,'original':dict(original)})
            return self._proposal(conn,actor,e,now,'legacy',{'treasury_id':p['treasury_id'],'rows':rows},
                                  {'treasury_id':p['treasury_id'],'rows':[{'table':r['table'],'id':r['id'],'operation_on':r['movement']['operation_on'],
                                   'amount_cents':r['movement']['amount_cents'],'quality':'observada'} for r in rows]})
        return self._write(session,env,'configure',op)

    def legacy_confirm(self,session,env):
        def op(conn,actor,e,now):
            self._session(conn,session,e.community_id,'import')
            proposal=self._row(conn,'erp_conciliacion_propuestas',e.community_id,e.payload['proposal_id']);self._version(proposal,e.expected_version)
            if proposal['kind']!='legacy' or proposal['state']!='pendiente':raise ConflictError('Propuesta historica no disponible.')
            data=self.vault.get(conn,e.community_id,proposal['secret_id'],'reconciliation-proposal')['payload']
            for r in data['rows']:
                if conn.execute('SELECT 1 FROM erp_banco_legacy_correspondencias WHERE id_comunidad=? AND legacy_table=? AND legacy_id=?',
                                (e.community_id,r['table'],r['id'])).fetchone():raise ConflictError('Correspondencia historica ya incorporada.')
            from dataclasses import replace
            child=self.__class__.in_transaction(self.database_path,conn,vault=self.vault)
            preview=child.statement_preview(session,replace(e,command='erp5.statement.preview',expected_version=None,
                idempotency_key=e.idempotency_key+'-stage',payload={'treasury_id':data['treasury_id'],'profile':'manual-v1','rows':[r['movement'] for r in data['rows']]}))['entity']
            confirmed=child.statement_confirm(session,replace(e,command='erp5.statement.confirm',expected_version=preview['version'],
                idempotency_key=e.idempotency_key+'-confirm',payload={'id':preview['id']}))['entity']
            if len(confirmed['movement_ids'])!=len(data['rows']):raise ConflictError('Correspondencias incompletas.')
            for r,mid in zip(data['rows'],confirmed['movement_ids']):
                conn.execute('INSERT INTO erp_banco_legacy_correspondencias (id_comunidad,legacy_table,legacy_id,movement_id,registered_at,actor_id) VALUES (?,?,?,?,?,?)',
                             (e.community_id,r['table'],r['id'],mid,now,actor.user_id))
            conn.execute("UPDATE erp_conciliacion_propuestas SET state='confirmada',version=version+1 WHERE id=?",(proposal['id'],))
            return {'id':proposal['id'],'movement_ids':confirmed['movement_ids'],'quality':'observada','activated':False}
        return self._write(session,env,'configure',op)

    def _report_summary(self,conn,session,community,p):
        from .receivables_projection import collection_balance
        free=None
        try:ReceivablesService._session(conn,session,community,'sensitive_read')
        except PermissionError:pass
        else:free=str(sum(int(collection_balance(conn,community,r['id'],p['end_on'])['available_cents']) for r in conn.execute(
            'SELECT id FROM erp_cobros WHERE id_comunidad=? AND treasury_reference=? AND effective_on<=?',(community,'erp4-treasury:'+str(p['treasury_id']),p['end_on']))))
        transit=0
        for t in conn.execute('SELECT * FROM erp_tesoreria_transferencias WHERE id_comunidad=? AND (source_id=? OR destination_id=?)',
            (community,p['treasury_id'],p['treasury_id'])):
            legs=list(conn.execute('SELECT * FROM erp_tesoreria_transferencia_extremos WHERE id_comunidad=? AND transfer_id=? AND effective_on<=?',(community,t['id'],p['end_on'])))
            if legs:transit+=t['amount_cents']-min(sum(abs(r['amount_cents']) for r in legs if r['side']==side) for side in ('source','destination'))
        return {'source':'ERP5','treasury_id':p['treasury_id'],'start_on':p['start_on'],'end_on':p['end_on'],
            'reported_cents':p['reported_cents'],'calculated_cents':p['calculated_cents'],'difference_cents':p['difference_cents'],
            'unidentified_cents':str(sum(abs(int(r['remaining_cents'])) for r in p['pending'])),
            'unapplied_funds_cents':free,'in_transit_cents':str(transit),'complete':p['complete']}

    def report_get(self,session,q):
        def op(conn,q):
            plan=self._balances(conn,q.community_id,q.filters)
            rows=[self._movement_public(conn,q.community_id,self._row(conn,'erp_banco_movimientos',q.community_id,mid)) for mid in plan['movement_ids']]
            return {'balances':plan,'summary':self._report_summary(conn,session,q.community_id,plan),'rows':rows,'source':'ERP5','account':plan['treasury_id'],'cutoff':plan['end_on'],
                    'economic_effect':False,'banking_data':'masked'}
        return self._read(session,q,'export',op)

    def report_export(self,session,env):
        def op(conn,actor,e,now):
            import base64,csv,hashlib,io
            require_fields(e.payload,('treasury_id','start_on','end_on'),('format',))
            p=self._balances(conn,e.community_id,e.payload);fmt=e.payload.get('format','csv')
            if fmt not in ('csv','xlsx','pdf'):raise ContractError('Elige CSV, Excel o PDF.')
            rows=[self._movement_public(conn,e.community_id,self._row(conn,'erp_banco_movimientos',e.community_id,mid)) for mid in p['movement_ids']]
            data=[['Fecha operacion','Fecha valor','Concepto','Importe EUR','Estado','Pendiente EUR']]
            def money(value):
                if value is None:return 'No verificable / no autorizado'
                value=int(value);n=abs(value)
                return ('-' if value<0 else '')+str(n//100)+'.'+str(n%100).zfill(2)
            def safe(value):
                value=str(value or '')
                return "'"+value if value.lstrip()[:1] in ('=','+','-','@') or value[:1] in ('\t','\r') else value
            data += [[r['operation_on'],r['value_on'] or 'No consta',safe(r['concept']),money(r['amount_cents']),r['state'],money(r['remaining_cents'])] for r in rows]
            summary=self._report_summary(conn,session,e.community_id,p)
            info=[['Fuente','ERP 5 - Banco; importes EUR'],['Periodo',p['start_on']+' - '+p['end_on']],
                ['Saldo comunicado',money(summary['reported_cents'])],['Saldo reconstruido',money(summary['calculated_cents'])],
                ['Diferencia',money(summary['difference_cents'])],['Sin identificar (magnitud)',money(summary['unidentified_cents'])],
                ['Fondos sin aplicar ERP 3',money(summary['unapplied_funds_cents'])],['Transferencias en transito',money(summary['in_transit_cents'])],
                ['Cobertura','Completa' if summary['complete'] else 'Sin acreditar']]
            if fmt=='csv':
                buffer=io.StringIO();writer=csv.writer(buffer,delimiter=';');writer.writerows(data+[[],*info]);raw=buffer.getvalue().encode('utf-8-sig')
            elif fmt=='pdf':
                from .reconciliation_reports import pdf_report
                raw=pdf_report(info,data)
            else:
                from openpyxl import Workbook
                from openpyxl.styles import Font
                book=Workbook();sheet=book.active;sheet.title='Conciliacion'
                for values in data:sheet.append(values)
                for cell in sheet[1]:cell.font=Font(bold=True)
                for key,width in (('A',18),('B',18),('C',60),('D',18),('E',22),('F',18)):sheet.column_dimensions[key].width=width
                sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
                summary_sheet=book.create_sheet('Resumen')
                for r in info:summary_sheet.append(r)
                summary_sheet.column_dimensions['A'].width=36;summary_sheet.column_dimensions['B'].width=50
                pending_sheet=book.create_sheet('Pendientes');pending_sheet.append(['Fecha','Concepto','Pendiente EUR','Responsable','Revisar el','Motivo'])
                for r in rows:
                    if not int(r['remaining_cents']):continue
                    note=conn.execute('SELECT * FROM erp_conciliacion_pendientes WHERE id_comunidad=? AND movement_id=? ORDER BY id DESC LIMIT 1',(e.community_id,r['id'])).fetchone()
                    user=conn.execute('SELECT nombre FROM usuarios WHERE id_usuario=?',(note['user_id'],)).fetchone() if note else None
                    content=self.vault.get(conn,e.community_id,note['secret_id'],'bank-pending')['note'] if note else 'Sin documentar'
                    from .banking_tabular import masked_cell
                    pending_sheet.append([r['operation_on'],safe(r['concept']),money(r['remaining_cents']),safe(masked_cell(user[0] if user else 'Sin asignar')),
                        note['review_on'] if note else '',safe(masked_cell(content))])
                for col,width in (('A',18),('B',60),('C',18),('D',25),('E',18),('F',60)):pending_sheet.column_dimensions[col].width=width
                pending_sheet.freeze_panes='A2';pending_sheet.auto_filter.ref=pending_sheet.dimensions
                buffer=io.BytesIO();book.save(buffer);raw=buffer.getvalue()
            digest=hashlib.sha256(raw).hexdigest()
            self._event(conn,e,'report.exported','bank-report',digest,{'effect':'evidence','format':fmt,'treasury_id':p['treasury_id'],
                'start_on':p['start_on'],'end_on':p['end_on'],'rows':len(rows),'sha256':digest})
            return {'filename':'conciliacion-'+p['end_on']+'.'+fmt,'content_base64':base64.b64encode(raw).decode('ascii'),
                    'sha256':digest,'banking_data':'masked','rows':len(rows)}
        return self._write(session,env,'export',op)

    def source_preview(self,session,env):
        def op(conn,actor,e,now):
            plan=self._closure_plan(conn,e.community_id,e.payload)
            if conn.execute('SELECT 1 FROM erp_banco_fuentes_activaciones WHERE id_comunidad=? AND treasury_id=? AND start_on<=? AND end_on>=?',
                            (e.community_id,plan['treasury_id'],plan['end_on'],plan['start_on'])).fetchone():
                raise ConflictError('Ya existe una fuente ERP 5 activa para este intervalo. No se activa otro origen paralelo.')
            return self._proposal(conn,actor,e,now,'source_activation',e.payload,plan)
        return self._write(session,env,'configure',op)

    def source_confirm(self,session,env):
        def op(conn,actor,e,now):
            proposal=self._row(conn,'erp_conciliacion_propuestas',e.community_id,e.payload['proposal_id']);self._version(proposal,e.expected_version)
            if proposal['state']!='pendiente' or proposal['kind']!='source_activation':raise ConflictError('Propuesta no disponible.')
            data=self.vault.get(conn,e.community_id,proposal['secret_id'],'reconciliation-proposal');plan=self._closure_plan(conn,e.community_id,data['payload'])
            if plan['preview_hash']!=data['public']['preview_hash']:raise ConflictError('Han cambiado cobertura o correspondencias. Revisa de nuevo.')
            if conn.execute('SELECT 1 FROM erp_banco_fuentes_activaciones WHERE id_comunidad=? AND treasury_id=? AND start_on<=? AND end_on>=?',
                            (e.community_id,plan['treasury_id'],plan['end_on'],plan['start_on'])).fetchone():raise ConflictError('Cobertura ya activada.')
            secret=self.vault.put(conn,e.community_id,'bank-source-activation',{'plan':plan,'evidence':e.evidence.entity_id,'reason':e.reason},now)
            rid=conn.execute('INSERT INTO erp_banco_fuentes_activaciones (id_comunidad,treasury_id,start_on,end_on,secret_id,registered_at,actor_id) VALUES (?,?,?,?,?,?,?)',
                            (e.community_id,plan['treasury_id'],plan['start_on'],plan['end_on'],secret,now,actor.user_id)).lastrowid
            conn.execute("UPDATE erp_conciliacion_propuestas SET state='confirmada',version=version+1 WHERE id=?",(proposal['id'],))
            self._event(conn,e,'source.activated','bank-source',rid,{'effect':'evidence','treasury_id':plan['treasury_id'],'start_on':plan['start_on'],'end_on':plan['end_on']})
            return {'id':rid,'source':'ERP5','economic_legacy_activated':False,'bank_statement_source_activated':True}
        return self._write(session,env,'configure',op)

    def source_status(self,session,q):
        def op(conn,q):
            p=q.filters;account=self._account(conn,q.community_id,p['treasury_id'])
            rows=[dict(r) for r in conn.execute('SELECT id,start_on,end_on,registered_at FROM erp_banco_fuentes_activaciones WHERE id_comunidad=? AND treasury_id=? AND start_on<=? AND end_on>=?',
                                               (q.community_id,account['id'],day(p['end_on']),day(p['start_on'])))]
            return {'items':rows,'source':'ERP5' if rows else 'ERP5 sin cobertura rectora activada','legacy_joined':False}
        return self._read(session,q,'read',op)
