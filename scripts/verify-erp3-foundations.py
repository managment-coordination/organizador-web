"""ERP 3 migration/transaction tests on an isolated SQLite backup, never production."""

import hashlib
import base64
import io
import json
from pathlib import Path
import sqlite3
import sys
import tempfile
import unittest
from contextlib import closing
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'server'))

from access_control import profile
from erp_core.contracts import CommandEnvelope, QueryEnvelope
from erp_core.database import connect
from erp_core.errors import ConflictError, ContractError, NotFoundError
from erp_core.migrations import apply_all
from erp_core.receivables_contracts import cents, day, known_time
from erp_core.receivables_projection import receipt_balance, collection_balance, responsibility_balance
from erp_core.receivables_service import ReceivablesService
from erp_core.dispatcher import execute_command, execute_query


SOURCE=Path(sys.argv[1]).resolve()
sys.argv=sys.argv[:1]
WORK=Path(tempfile.mkdtemp(prefix='organizador-erp3-foundations-'))


def backup(source,target):
    with closing(sqlite3.connect('file:'+str(source)+'?mode=ro',uri=True)) as src:
        with closing(sqlite3.connect(target)) as dst:src.backup(dst)


class ERP3Tests(unittest.TestCase):
    def setUp(self):
        self.db=WORK/(self._testMethodName+'.db');backup(SOURCE,self.db)
        self.conn=connect(self.db)
        self.before={t:self.conn.execute('SELECT count(*) FROM '+t).fetchone()[0] for t in ('cf_recibos','cf_movimientos_deuda','cf_propietarios','cf_propiedades')}
        apply_all(self.conn)
        self.uid=self.conn.execute("SELECT id_usuario FROM usuarios WHERE rol='Superusuario' AND activo=1 LIMIT 1").fetchone()[0]
        self.session=profile(self.conn,self.uid)
        self.community=self.session['comunidades'][0]['id_comunidad']
        self.property=self.conn.execute('SELECT id_propiedad FROM cf_propiedades WHERE id_comunidad=? LIMIT 1',(self.community,)).fetchone()[0]
        self.conn.execute('''INSERT INTO erp_ejercicios (id_comunidad,codigo,fecha_inicio,fecha_fin,estado,creado_en,creado_por,origen)
            VALUES (?,'ERP3-TEST','2026-01-01','2026-12-31','abierto','2026-01-01',?,'test')''',(self.community,self.uid))
        self.exercise=self.conn.execute('SELECT last_insert_rowid()').fetchone()[0]
        self.service=ReceivablesService(self.db)
        self.n=0

    def tearDown(self):
        for t,count in self.before.items():self.assertEqual(self.conn.execute('SELECT count(*) FROM '+t).fetchone()[0],count)
        self.assertEqual(self.conn.execute('PRAGMA integrity_check').fetchone()[0],'ok')
        self.assertEqual(self.conn.execute('PRAGMA foreign_key_check').fetchall(),[])
        self.conn.close()

    def command(self,name,payload,version=None,key=None,session=None,evidence=None):
        self.n+=1
        value={'command':'erp3.'+name,'id_comunidad':self.community,'payload':payload,
            'idempotency_key':key or 'test-'+str(self.n),'expected_version':version,'reason':'Verificacion sintetica','origin':'test','evidence':evidence}
        return execute_command(self.db,session or self.session,value)

    def receipt(self,amount=10000,number='TEST-1'):
        return self.conn.execute('''INSERT INTO erp_recibos
            (id_comunidad,id_propiedad,id_ejercicio,number,concept_key,description,period_key,period_from,period_until,
             source_type,source_key,obligation_key,amount_cents,currency,issued_on,due_on,snapshot_json,snapshot_hash,registered_at,actor_id)
             VALUES (?,?,?,?,'test','Cuota sintetica','P01','2026-01-01','2026-01-31','plan',?,?,?,'EUR',
             '2026-01-01','2026-01-10','{}','test','2026-01-01T00:00:00.000000Z',?)''',
             (self.community,self.property,self.exercise,number,number,number,amount,self.uid)).lastrowid

    def collection(self,amount=10000,key='BANK-1'):
        return self.command('collection.record',{'amount_cents':str(amount),'currency':'EUR','effective_on':'2026-01-15',
            'method':'transferencia','external_source':'synthetic','external_key':key})['entity']['id']

    def allocate(self,cid,rid,amount):
        proposal=self.command('allocation.preview',{'collection_id':cid,'effective_on':'2026-01-15',
            'allocations':[{'receipt_id':rid,'amount_cents':str(amount)}]})['entity']
        return self.command('allocation.confirm',{'proposal_id':proposal['id']},proposal['version'])

    def test_migration_reentrant_restore(self):
        checks=list(self.conn.execute('SELECT version,checksum FROM erp_schema_migrations'))
        apply_all(self.conn)
        self.assertEqual([tuple(r) for r in checks],[tuple(r) for r in self.conn.execute('SELECT version,checksum FROM erp_schema_migrations')])
        restored=WORK/'restored.db';backup(self.db,restored)
        with closing(sqlite3.connect(restored)) as conn:
            self.assertEqual(conn.execute('PRAGMA integrity_check').fetchone()[0],'ok')
            self.assertEqual(conn.execute('PRAGMA foreign_key_check').fetchall(),[])

    def test_exact_validation(self):
        self.assertEqual(cents('100'),100)
        for bad in (1.2,True,'1.00','1e3',None,9000000000000001):
            with self.assertRaises(ContractError):cents(bad)
        with self.assertRaises(ContractError):day('20260101')
        with self.assertRaises(ContractError):known_time('2026-01-01T00:00:00')

    def test_paid_credit_frees_funds_atomically(self):
        rid=self.receipt();cid=self.collection();self.allocate(cid,rid,10000)
        aid=self.conn.execute('SELECT id FROM erp_imputaciones WHERE receipt_id=?',(rid,)).fetchone()[0]
        with self.assertRaises(ConflictError):
            self.command('credit.preview',{'receipt_id':rid,'effective_on':'2026-01-16','amount_cents':'2000'})
        p=self.command('credit.preview',{'receipt_id':rid,'effective_on':'2026-01-16','amount_cents':'2000',
            'release_allocations':[{'allocation_id':aid,'amount_cents':'2000'}]})['entity']
        self.command('credit.confirm',{'proposal_id':p['id']},1)
        self.assertEqual(receipt_balance(self.conn,self.community,rid)['pending_cents'],'0')
        self.assertEqual(receipt_balance(self.conn,self.community,rid)['paid_cents'],'8000')
        self.assertEqual(collection_balance(self.conn,self.community,cid)['available_cents'],'2000')
        self.assertEqual(self.conn.execute('SELECT amount_cents FROM erp_recibos WHERE id=?',(rid,)).fetchone()[0],10000)

    def test_desimputation_without_return_and_replay(self):
        rid=self.receipt();cid=self.collection();self.allocate(cid,rid,10000)
        aid=self.conn.execute('SELECT id FROM erp_imputaciones WHERE receipt_id=?',(rid,)).fetchone()[0]
        p=self.command('allocation.reverse.preview',{'effective_on':'2026-01-17','reversals':[{'allocation_id':aid,'amount_cents':'3000'}]})['entity']
        self.command('allocation.reverse.confirm',{'proposal_id':p['id']},1,key='reverse-confirm')
        replay=self.command('allocation.reverse.confirm',{'proposal_id':p['id']},1,key='reverse-confirm')
        self.assertTrue(replay['idempotent_replay'])
        self.assertEqual(receipt_balance(self.conn,self.community,rid)['pending_cents'],'3000')
        self.assertEqual(collection_balance(self.conn,self.community,cid)['available_cents'],'3000')
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_devoluciones').fetchone()[0],0)

    def test_return_fee_policies_separate_charge(self):
        rid=self.receipt();cid=self.collection()
        owner=self.conn.execute('SELECT id_propietario,nombre FROM cf_propietarios WHERE id_comunidad=? LIMIT 1',(self.community,)).fetchone()
        self.conn.execute("INSERT INTO erp_recibo_sujetos (id_comunidad,receipt_id,role,owner_id,snapshot_json) VALUES (?,?,'obligated',?,?)",
            (self.community,rid,owner[0],json.dumps({'type':'owner','id':owner[0],'name':owner[1]})))
        self.conn.execute('''INSERT INTO erp_config_recibo_versiones
            (id_comunidad,id_propiedad,version,efectiva_desde,registrada_en,registrada_por,estado,destinatario_propietario_id,pagador_propietario_id,medio_previsto,origen)
            VALUES (?,?,1000,'2026-01-01','2026-01-01',?,'confirmada',?,?,'transferencia','test')''',(self.community,self.property,self.uid,owner[0],owner[0]))
        self.allocate(cid,rid,10000)
        aid=self.conn.execute('SELECT id FROM erp_imputaciones WHERE receipt_id=?',(rid,)).fetchone()[0]
        p=self.command('return.preview',{'collection_id':cid,'effective_on':'2026-01-17','free_cents':'0',
            'reversals':[{'allocation_id':aid,'amount_cents':'10000'}],'external_key':'RETURN-COST'})['entity']
        self.command('return.confirm',{'proposal_id':p['id']},1)
        did=self.conn.execute('SELECT id FROM erp_devoluciones').fetchone()[0]
        for version,mode,expected in [(0,'none',0),(1,'actual',500),(2,'fixed',700)]:
            self.command('policy.save',{'effective_from':'2026-01-17','return_fee_mode':mode,'fixed_cents':'700' if mode=='fixed' else '0'},version)
            p=self.command('return.fee.preview',{'return_id':did,'receipt_id':rid,'effective_on':'2026-01-18',
                'cost_cents':'500','cost_reference':'COST-'+mode,'exercise_id':self.exercise})['entity']
            self.assertEqual(p['preview']['amount_cents'],str(expected))
            result=self.command('return.fee.confirm',{'proposal_id':p['id']},1,evidence={'type':'external_reference','id':'synthetic:cost'})['entity']
            if expected:self.assertEqual(receipt_balance(self.conn,self.community,result['receipt_id'])['pending_cents'],str(expected))
            else:self.assertIsNone(result['receipt_id'])
        self.assertEqual(receipt_balance(self.conn,self.community,rid)['original_cents'],'10000')
        with self.assertRaises(ConflictError):
            self.command('return.fee.preview',{'return_id':did,'receipt_id':rid,'effective_on':'2026-01-18','cost_cents':'500','cost_reference':'COST-fixed'})

    def test_refund_excess_does_not_reopen_receipt(self):
        rid=self.receipt();owner=self.conn.execute('SELECT id_propietario FROM cf_propietarios WHERE id_comunidad=? LIMIT 1',(self.community,)).fetchone()[0]
        cid=self.command('collection.record',{'amount_cents':'12000','currency':'EUR','effective_on':'2026-01-15',
            'method':'transferencia','external_source':'test','external_key':'REFUND','payer':{'type':'owner','id':owner}})['entity']['id']
        self.allocate(cid,rid,10000)
        p=self.command('refund.preview',{'collection_id':cid,'amount_cents':'2000','effective_on':'2026-01-20',
            'beneficiary':{'type':'owner','id':owner}})['entity']
        self.command('refund.confirm',{'proposal_id':p['id']},1,evidence={'type':'external_reference','id':'synthetic:refund'})
        self.assertEqual(receipt_balance(self.conn,self.community,rid)['pending_cents'],'0')
        self.assertEqual(collection_balance(self.conn,self.community,cid)['available_cents'],'0')
        self.assertEqual(collection_balance(self.conn,self.community,cid)['refunded_cents'],'2000')

    def test_exceptional_transfer_and_explicit_later_payment(self):
        rid=self.receipt();owners=[r[0] for r in self.conn.execute('SELECT id_propietario FROM cf_propietarios WHERE id_comunidad=? LIMIT 2',(self.community,))]
        a=[{'type':'owner','id':owners[0]}];b=[{'type':'owner','id':owners[1]}]
        self.conn.execute("INSERT INTO erp_recibo_sujetos (id_comunidad,receipt_id,role,owner_id,snapshot_json) VALUES (?,?,'obligated',?,?)",
            (self.community,rid,owners[0],json.dumps(a[0])))
        original=self.conn.execute('SELECT snapshot_hash FROM erp_recibos WHERE id=?',(rid,)).fetchone()[0]
        p=self.command('responsibility.transfer.preview',{'effective_on':'2026-01-10','authorization_reference':'Acuerdo sintetico',
            'lines':[{'receipt_id':rid,'amount_cents':'4000','source_subjects':a,'target_subjects':b}]})['entity']
        self.command('responsibility.transfer.confirm',{'proposal_id':p['id']},1,key='transfer',evidence={'type':'external_reference','id':'synthetic:transfer'})
        self.assertEqual(sorted(int(x['pending_cents']) for x in responsibility_balance(self.conn,self.community,rid)),[4000,6000])
        self.assertEqual(responsibility_balance(self.conn,self.community,rid,'2026-01-09')[0]['pending_cents'],'10000')
        self.assertEqual(responsibility_balance(self.conn,self.community,rid,'2026-01-20','2026-02-01T00:00:00Z')[0]['pending_cents'],'10000')
        cid=self.collection(3000)
        with self.assertRaises(ConflictError):self.allocate(cid,rid,3000)
        p=self.command('allocation.preview',{'collection_id':cid,'effective_on':'2026-01-15','allocations':[
            {'receipt_id':rid,'amount_cents':'3000','responsibility_subjects':b}]})['entity']
        self.command('allocation.confirm',{'proposal_id':p['id']},1)
        balances=responsibility_balance(self.conn,self.community,rid)
        self.assertEqual(next(x['pending_cents'] for x in balances if x['subjects']==a),'6000')
        self.assertEqual(next(x['pending_cents'] for x in balances if x['subjects']==b),'1000')
        aid=self.conn.execute('SELECT id FROM erp_imputaciones WHERE receipt_id=?',(rid,)).fetchone()[0]
        p=self.command('return.preview',{'collection_id':cid,'effective_on':'2026-01-18','free_cents':'0','external_key':'TRANSFER-RETURN',
            'reversals':[{'allocation_id':aid,'amount_cents':'3000'}]})['entity']
        self.command('return.confirm',{'proposal_id':p['id']},1)
        self.assertEqual(sorted(int(x['pending_cents']) for x in responsibility_balance(self.conn,self.community,rid)),[4000,6000])
        self.assertEqual(self.conn.execute('SELECT snapshot_hash FROM erp_recibos WHERE id=?',(rid,)).fetchone()[0],original)

    def test_history_staging_cutoff_identity_and_rollback(self):
        row={'reference':'OPENING-2024','amount_cents':'530000','cutoff_date':'2024-12-31','coverage_from':'2024-01-01',
             'coverage_until':'2024-12-31','scope':'receipt','property_id':self.property,'limitations':'Saldo observado; cobros no reconstruibles.'}
        payload={'source':'synthetic-history','file_hash':'a'*64,'file_name':'historico.xlsx','rows':[row]}
        p=self.command('history.import.preview',payload)['entity']
        self.assertEqual(p['blocking_issues'],0)
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM erp_saldos_apertura').fetchone()[0],0)
        self.command('history.import.confirm',{'import_id':p['id']},1,evidence={'type':'external_reference','id':'synthetic:history'})
        self.assertEqual(self.conn.execute('SELECT amount_cents,effective_on,quality FROM erp_saldos_apertura').fetchone()[:],(530000,'2024-12-31','observada'))
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM erp_cobros').fetchone()[0],0)
        replay=self.command('history.import.preview',payload)['entity'];self.assertTrue(replay['duplicate'])
        other=self.command('history.import.preview',{**payload,'file_hash':'b'*64})['entity']
        self.assertEqual(other['rows'][0]['decision'],'skip')
        self.command('history.import.confirm',{'import_id':other['id']},1,evidence={'type':'external_reference','id':'synthetic:history'})
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM erp_saldos_apertura').fetchone()[0],1)
        conflict=self.command('history.import.preview',{**payload,'file_hash':'c'*64,'rows':[{**row,'amount_cents':'1'}]})['entity']
        self.assertEqual(conflict['blocking_issues'],1)
        with self.assertRaises(ConflictError):self.command('history.import.confirm',{'import_id':conflict['id']},1,evidence={'type':'external_reference','id':'synthetic:history'})
        with self.assertRaises(sqlite3.IntegrityError):self.conn.execute("UPDATE erp_saldos_apertura SET amount_cents=1")

    def activation_fixture(self):
        rid=self.conn.execute('''INSERT INTO cf_recibos (referencia,fecha_emision,id_propiedad,id_comunidad,importe,cobrado,deuda)
            VALUES ('SYNTHETIC-ACTIVATION','2030-01-03',?,?,100,25,75)''',(self.property,self.community)).lastrowid
        self.before['cf_recibos']+=1
        row={'reference':'ACTIVATION-SALDO','amount_cents':'7500','cutoff_date':'2030-01-31','coverage_from':'2030-01-01',
             'coverage_until':'2030-01-31','scope':'receipt','property_id':self.property,'legacy_receipt_ids':[rid],
             'limitations':'Observacion sintetica, no se reconstruyen pagos.'}
        payload={'source':'activation-test','file_hash':'d'*64,'file_name':'historico.xlsx','rows':[row]}
        draft=self.command('history.import.preview',payload)['entity']
        opened=self.command('history.import.confirm',{'import_id':draft['id']},draft['version'],evidence={'type':'external_reference','id':'synthetic-history'})['entity']
        coverage=self.command('coverage.confirm',{'concept_key':'ordinario','effective_from':'2030-01-01','effective_until':'2030-01-31','authority':'legacy_observed'},
                              evidence={'type':'external_reference','id':'synthetic-cut'})['entity']['id']
        return rid,payload,{'coverage_id':coverage,'mappings':[{'legacy_receipt_id':rid,'opening_id':opened['opening_ids'][0],
                 'period_from':'2030-01-01','period_until':'2030-01-31','concept_key':'ordinario'}]}

    def test_historical_activation_review_idempotency_and_originals(self):
        rid,imported,payload=self.activation_fixture()
        original=tuple(self.conn.execute('SELECT * FROM cf_recibos WHERE id_recibo=?',(rid,)).fetchone())
        with self.assertRaises(ConflictError):self.service._require_coverage(self.conn,self.community,'ordinario','2030-01-01','2030-01-31')
        with self.assertRaises(ConflictError):self.command('coverage.activation.preview',{**payload,'mappings':[]})
        preview=self.command('coverage.activation.preview',payload)['entity']
        self.assertEqual(preview['preview']['lines'][0]['net_emitted_cents'],'10000')
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM erp_historico_obligaciones').fetchone()[0],0)
        result=self.command('coverage.activation.confirm',{'proposal_id':preview['id']},preview['version'],key='activate',evidence={'type':'external_reference','id':'reviewed-correspondence'})
        self.assertTrue(self.command('coverage.activation.confirm',{'proposal_id':preview['id']},preview['version'],key='activate',evidence={'type':'external_reference','id':'reviewed-correspondence'})['idempotent_replay'])
        self.assertFalse(result['entity']['receipts_created'])
        self.service._require_coverage(self.conn,self.community,'ordinario','2030-01-01','2030-01-31')
        self.assertEqual(tuple(self.conn.execute('SELECT * FROM cf_recibos WHERE id_recibo=?',(rid,)).fetchone()),original)
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM erp_cobros').fetchone()[0],0)
        self.assertEqual(self.conn.execute('SELECT quality FROM erp_saldos_apertura').fetchone()[0],'observada')
        with self.assertRaises(sqlite3.IntegrityError):self.conn.execute('UPDATE erp_historico_obligaciones SET net_emitted_cents=0')
        self.assertEqual(self.command('history.import.preview',{**imported,'file_hash':'e'*64})['entity']['rows'][0]['decision'],'skip')
        changed={**imported,'file_hash':'f'*64,'rows':[{**imported['rows'][0],'reference':'EXTRA','legacy_receipt_ids':[]}]}
        self.assertEqual(self.command('history.import.preview',changed)['entity']['blocking_issues'],1)

    def test_historical_activation_atomic_rollback(self):
        rid,_,payload=self.activation_fixture()
        preview=self.command('coverage.activation.preview',payload)['entity']
        self.conn.execute("CREATE TRIGGER synthetic_activation_failure BEFORE INSERT ON erp_historico_obligaciones BEGIN SELECT RAISE(ABORT,'synthetic failure'); END")
        with self.assertRaises(Exception):self.command('coverage.activation.confirm',{'proposal_id':preview['id']},preview['version'],evidence={'type':'external_reference','id':'reviewed-correspondence'})
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM erp_cobertura_activaciones').fetchone()[0],0)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM erp_hechos_economicos WHERE event_type='erp3.coverage.activated'").fetchone()[0],0)
        self.conn.execute('DROP TRIGGER synthetic_activation_failure')
        self.command('coverage.activation.confirm',{'proposal_id':preview['id']},preview['version'],evidence={'type':'external_reference','id':'reviewed-correspondence'})

    def test_batch_allocation_many_to_many_atomic_and_replay(self):
        r1=self.receipt(10000,'BATCH-1');r2=self.receipt(10000,'BATCH-2')
        c1=self.collection(10000,'BATCH-C1');c2=self.collection(10000,'BATCH-C2')
        payload={'effective_on':'2026-01-15','collections':[
            {'collection_id':c1,'allocations':[{'receipt_id':r1,'amount_cents':'6000'},{'receipt_id':r2,'amount_cents':'4000'}]},
            {'collection_id':c2,'allocations':[{'receipt_id':r1,'amount_cents':'4000'},{'receipt_id':r2,'amount_cents':'6000'}]}]}
        p=self.command('allocation.batch.preview',payload)['entity']
        self.assertEqual(p['preview']['total_cents'],'20000')
        self.conn.execute(f"CREATE TRIGGER synthetic_batch_failure BEFORE INSERT ON erp_imputaciones WHEN NEW.collection_id={c2} BEGIN SELECT RAISE(ABORT,'synthetic failure'); END")
        with self.assertRaises(Exception):self.command('allocation.batch.confirm',{'proposal_id':p['id']},p['version'],key='batch-confirm')
        self.assertEqual(collection_balance(self.conn,self.community,c1)['available_cents'],'10000')
        self.assertEqual(receipt_balance(self.conn,self.community,r1)['pending_cents'],'10000')
        self.conn.execute('DROP TRIGGER synthetic_batch_failure')
        self.command('allocation.batch.confirm',{'proposal_id':p['id']},p['version'],key='batch-confirm')
        self.assertTrue(self.command('allocation.batch.confirm',{'proposal_id':p['id']},p['version'],key='batch-confirm')['idempotent_replay'])
        self.assertEqual(receipt_balance(self.conn,self.community,r1)['pending_cents'],'0')
        self.assertEqual(receipt_balance(self.conn,self.community,r2)['pending_cents'],'0')

    def test_batch_rejects_combined_overapplication(self):
        rid=self.receipt();a=self.collection(10000,'A');b=self.collection(10000,'B')
        with self.assertRaises(ConflictError):self.command('allocation.batch.preview',{'effective_on':'2026-01-15',
            'collections':[{'collection_id':c,'allocations':[{'receipt_id':rid,'amount_cents':'6000'}]} for c in (a,b)]})
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM erp_imputaciones').fetchone()[0],0)

    def test_export_all_pages_exact_audited_formula_safe(self):
        from openpyxl import load_workbook
        for i in range(51):self.receipt(101,'=UNTRUSTED-'+str(i))
        payload={'kind':'receipts','format':'csv','filters':{'property_id':self.property,'effective_at':'2026-01-31','limit':1}}
        result=self.command('export.prepare',payload,key='export-request')
        self.assertEqual(result['entity']['row_count'],51)
        self.assertEqual(result['entity']['subtotal_cents'],'5151')
        content=base64.b64decode(result['entity']['content_base64']).decode('utf-8-sig')
        self.assertIn("'=UNTRUSTED-50",content)
        self.assertIn('51.51',content)
        self.assertTrue(self.command('export.prepare',payload,key='export-request')['idempotent_replay'])
        xlsx=self.command('export.prepare',{**payload,'format':'xlsx'})['entity']
        book=load_workbook(io.BytesIO(base64.b64decode(xlsx['content_base64'])))
        self.assertEqual(book['Datos'].max_row,52)
        self.assertTrue(all(cell.data_type!='f' for row in book['Datos'] for cell in row))
        self.assertTrue(book['Origen y alcance']['B2'].value=='2026-01-31')
        statement=execute_query(self.db,self.session,{'query':'erp3.account.statement','id_comunidad':self.community,
            'filters':{'effective_at':'2026-02-15','property_id':self.property}})['entity']
        self.assertEqual(statement['aging_cents']['31_60'],'5151')

    def test_pagination_exact_totals_and_personal_attribution(self):
        owners=[r[0] for r in self.conn.execute('SELECT id_propietario FROM cf_propietarios WHERE id_comunidad=? LIMIT 2',(self.community,))]
        for number,owner in [('DEBT-A',owners[0]),('DEBT-B',owners[1])]:
            rid=self.receipt(10000,number)
            self.conn.execute("INSERT INTO erp_recibo_sujetos (id_comunidad,receipt_id,role,owner_id,snapshot_json) VALUES (?,?,'obligated',?,?)",
                (self.community,rid,owner,json.dumps({'type':'owner','id':owner})))
        result=execute_query(self.db,self.session,{'query':'erp3.receipt.list','id_comunidad':self.community,
            'filters':{'property_id':self.property,'limit':1}})['entity']
        self.assertEqual(len(result['items']),1);self.assertEqual(result['total_count'],2)
        self.assertEqual(result['pending_cents'],'20000')
        summary=execute_query(self.db,self.session,{'query':'erp3.debt.summary','id_comunidad':self.community,
            'filters':{'owner_id':owners[0]}})['entity']
        self.assertEqual(summary['personal_pending_cents'],'10000')
        self.assertEqual(summary['shared_obligations_cents'],'0')
        self.assertIsNone(summary['total_pending_cents'])
        self.assertEqual(summary['observed_legacy_debt_cents'],'0')

    def test_evidence_scope_and_missing_document(self):
        with self.assertRaises(NotFoundError):
            self.command('policy.save',{'effective_from':'2026-01-01','return_fee_mode':'none'},0,
                evidence={'type':'attachment','id':'999999999'})
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_recibo_politicas').fetchone()[0],0)

    def test_financial_user_document_evidence_and_export_permissions(self):
        from access_control import permission
        candidates=[profile(self.conn,r[0]) for r in self.conn.execute("SELECT id_usuario FROM usuarios WHERE activo=1 AND rol<>'Superusuario'")]
        user=next(u for u in candidates if permission(u,self.community,'puede_ver_documentos'))
        self.command('permissions.save',{'user_id':user['id_usuario'],'capabilities':{'read':True,'sensitive_read':True,'credit':True}},0)
        doc=self.conn.execute("INSERT INTO anexos_registros(id_comunidad,tipo_entidad,id_registro,nombre_archivo,ruta_archivo,fecha_adjuntado) VALUES (?,'proyecto',0,'Justificante sintetico.pdf','synthetic-not-a-real-file','2026-01-01')",(self.community,)).lastrowid
        documents=execute_query(self.db,user,{'query':'erp3.evidence.list','id_comunidad':self.community,'filters':{}})['entity']
        self.assertTrue(any(d['id']==doc and d['type']=='attachment' for d in documents['items']))
        rid=self.receipt()
        p=self.command('credit.preview',{'receipt_id':rid,'effective_on':'2026-01-20','amount_cents':'1000'},session=user)['entity']
        self.command('credit.confirm',{'proposal_id':p['id']},p['version'],session=user,evidence={'type':'attachment','id':str(doc)})
        recorded=json.loads(self.conn.execute("SELECT evidence_json FROM erp_hechos_economicos WHERE event_type='erp3.credit.issued'").fetchone()[0])
        self.assertEqual(recorded,{'type':'attachment','id':str(doc)})
        export={'kind':'receipts','format':'csv','filters':{}}
        self.command('export.prepare',export,session=user)
        self.command('permissions.save',{'user_id':user['id_usuario'],'capabilities':{'read':True}},1)
        basic=execute_query(self.db,user,{'query':'erp3.receipt.list','id_comunidad':self.community,'filters':{}})['entity']
        self.assertTrue(all(not r['subjects'] and not r['responsibilities'] for r in basic['items']))
        with self.assertRaises(PermissionError):execute_query(self.db,user,{'query':'erp3.receipt.get','id_comunidad':self.community,'filters':{'receipt_id':rid}})
        with self.assertRaises(PermissionError):self.command('export.prepare',export,session=user)
        with self.assertRaises(PermissionError):self.command('coverage.activation.preview',{'coverage_id':1,'mappings':[]},session=user)

    def test_outbox_repeated_delivery_has_one_synthetic_consumer_effect(self):
        payload={'amount_cents':'10000','currency':'EUR','effective_on':'2026-01-15','method':'transferencia','external_source':'synthetic','external_key':'EVENT-ONCE'}
        first=self.command('collection.record',payload,key='event-once')
        second=self.command('collection.record',payload,key='event-once')
        self.assertEqual(first['outbox_event_id'],second['outbox_event_id'])
        event=self.conn.execute('SELECT * FROM erp_outbox WHERE event_id=?',(first['outbox_event_id'],)).fetchone()
        message=json.loads(event['payload_json'])
        self.assertEqual(len(message['economic_event_ids']),1)
        self.conn.execute('CREATE TEMP TABLE synthetic_consumer (community INTEGER,event_id INTEGER,amount_cents INTEGER,UNIQUE(community,event_id))')
        for _ in range(2):
            for eid in message['economic_event_ids']:
                fact=self.conn.execute('SELECT * FROM erp_hechos_economicos WHERE id_comunidad=? AND id=?',(self.community,eid)).fetchone()
                body=json.loads(fact['payload_json'])
                self.conn.execute('INSERT OR IGNORE INTO synthetic_consumer VALUES (?,?,?)',(self.community,eid,int(body['amount_cents'])))
        self.assertEqual(tuple(self.conn.execute('SELECT COUNT(*),SUM(amount_cents) FROM synthetic_consumer').fetchone()),(1,10000))

    def test_person_is_not_inferred_from_owner_or_payer(self):
        person=self.conn.execute("INSERT INTO erp_personas_cobro(id_comunidad,tipo,nombre,creada_en,creada_por,origen) VALUES (?,'fisica','Persona sintetica','2026-01-01',?,'test')",(self.community,self.uid)).lastrowid
        rid=self.receipt()
        self.conn.execute("INSERT INTO erp_recibo_sujetos (id_comunidad,receipt_id,role,person_id,snapshot_json) VALUES (?,?,'obligated',?,?)",(self.community,rid,person,json.dumps({'type':'person','id':person,'name':'Persona sintetica'})))
        result=execute_query(self.db,self.session,{'query':'erp3.account.statement','id_comunidad':self.community,'filters':{'person_id':person}})['entity']
        self.assertEqual(result['personal_cents'],'10000')
        self.assertEqual(len(result['items']),1)

    def test_shared_obligated_group_never_splits_or_doubles_charge(self):
        rid=self.receipt()
        owners=[r[0] for r in self.conn.execute('SELECT id_propietario FROM cf_propietarios WHERE id_comunidad=? LIMIT 2',(self.community,))]
        for oid in owners:self.conn.execute("INSERT INTO erp_recibo_sujetos(id_comunidad,receipt_id,role,owner_id,snapshot_json) VALUES (?,?,'obligated',?,?)",(self.community,rid,oid,json.dumps({'type':'owner','id':oid})))
        for oid in owners:
            data=execute_query(self.db,self.session,{'query':'erp3.debt.summary','id_comunidad':self.community,'filters':{'owner_id':oid}})['entity']
            self.assertEqual(data['shared_obligations_cents'],'10000');self.assertEqual(data['personal_pending_cents'],'0')
        self.assertEqual(self.conn.execute('SELECT COUNT(*),SUM(amount_cents) FROM erp_recibos').fetchone()[:],(1,10000))

    def test_uncollectible_discard_and_period_lock_preserve_debt(self):
        rid=self.receipt()
        p=self.command('uncollectible.preview',{'receipt_id':rid,'effective_on':'2026-01-20','classification':'incobrable'},evidence={'type':'external_reference','id':'SYNTHETIC-REASON'})['entity']
        self.command('uncollectible.confirm',{'proposal_id':p['id']},p['version'],evidence={'type':'external_reference','id':'SYNTHETIC-REASON'})
        self.assertEqual(receipt_balance(self.conn,self.community,rid)['pending_cents'],'10000')
        self.assertEqual(receipt_balance(self.conn,self.community,rid)['management'],'incobrable')
        p=self.command('credit.preview',{'receipt_id':rid,'effective_on':'2026-01-21','amount_cents':'100'})['entity']
        before=self.conn.execute('SELECT COUNT(*) FROM erp_hechos_economicos').fetchone()[0]
        self.command('proposal.discard',{'proposal_id':p['id']},p['version'])
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM erp_hechos_economicos').fetchone()[0],before)
        with self.assertRaises(ConflictError):self.command('credit.confirm',{'proposal_id':p['id']},p['version'])
        self.conn.execute("UPDATE erp_ejercicios SET estado='cerrado' WHERE id_ejercicio=?",(self.exercise,))
        with self.assertRaises(ConflictError):self.command('credit.preview',{'receipt_id':rid,'effective_on':'2026-01-22','amount_cents':'100'})

    def test_mixed_return_and_known_time_no_automatic_netting(self):
        rid=self.receipt(15000);cid=self.collection(20000);self.allocate(cid,rid,15000)
        aid=self.conn.execute('SELECT id FROM erp_imputaciones WHERE receipt_id=?',(rid,)).fetchone()[0]
        known_before=known_time(None)
        p=self.command('return.preview',{'collection_id':cid,'effective_on':'2026-01-20','free_cents':'5000','external_key':'MIXED-80',
            'reversals':[{'allocation_id':aid,'amount_cents':'3000'}]})['entity']
        self.command('return.confirm',{'proposal_id':p['id']},p['version'])
        self.assertEqual(receipt_balance(self.conn,self.community,rid)['pending_cents'],'3000')
        self.assertEqual(collection_balance(self.conn,self.community,cid)['available_cents'],'0')
        self.assertEqual(receipt_balance(self.conn,self.community,rid,'2026-01-20',known_before)['pending_cents'],'0')
        self.assertEqual(receipt_balance(self.conn,self.community,rid,'2026-01-14')['pending_cents'],'15000')
        p=self.command('return.preview',{'collection_id':cid,'effective_on':'2026-01-21','free_cents':'0','external_key':'REST',
            'reversals':[{'allocation_id':aid,'amount_cents':'12000'}]})['entity']
        self.command('return.confirm',{'proposal_id':p['id']},p['version'])
        with self.assertRaises(ConflictError):self.command('void.preview',{'receipt_id':rid,'effective_on':'2026-01-22'})

    def test_concurrent_confirmation_cannot_spend_twice(self):
        a=self.receipt();b=self.receipt(10000,'CONCURRENT');cid=self.collection()
        proposals=[self.command('allocation.preview',{'collection_id':cid,'effective_on':'2026-01-15',
            'allocations':[{'receipt_id':rid,'amount_cents':'10000'}]})['entity'] for rid in (a,b)]
        barrier=Barrier(2)
        def confirm(i):
            barrier.wait()
            try:
                execute_command(self.db,self.session,{'command':'erp3.allocation.confirm','id_comunidad':self.community,
                    'payload':{'proposal_id':proposals[i]['id']},'expected_version':1,'idempotency_key':'parallel-'+str(i),
                    'origin':'test','reason':'Confirmacion humana sintetica concurrente'})
                return 'confirmed'
            except ConflictError:return 'conflict'
        with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(confirm,(0,1)))
        self.assertEqual(sorted(results),['confirmed','conflict'])
        self.assertEqual(collection_balance(self.conn,self.community,cid)['available_cents'],'0')
        self.assertEqual(sum(int(receipt_balance(self.conn,self.community,r)['pending_cents']) for r in (a,b)),10000)

    def test_transfer_lot_failure_rolls_back_all_lines(self):
        owners=[r[0] for r in self.conn.execute('SELECT id_propietario FROM cf_propietarios WHERE id_comunidad=? LIMIT 2',(self.community,))]
        a=[{'type':'owner','id':owners[0]}];b=[{'type':'owner','id':owners[1]}]
        receipts=[self.receipt(10000,'TRANSFER-'+str(i)) for i in range(2)]
        for rid in receipts:self.conn.execute("INSERT INTO erp_recibo_sujetos (id_comunidad,receipt_id,role,owner_id,snapshot_json) VALUES (?,?,'obligated',?,?)",(self.community,rid,owners[0],json.dumps(a[0])))
        proposal=self.command('responsibility.transfer.preview',{'effective_on':'2026-01-10','authorization_reference':'Acuerdo sintetico',
            'lines':[{'receipt_id':r,'amount_cents':'4000','source_subjects':a,'target_subjects':b} for r in receipts]})['entity']
        self.conn.execute(f"CREATE TRIGGER fail_transfer BEFORE INSERT ON erp_reasignaciones_obligacion WHEN NEW.receipt_id={receipts[1]} BEGIN SELECT RAISE(ABORT,'synthetic fault'); END")
        events=self.conn.execute('SELECT count(*) FROM erp_hechos_economicos').fetchone()[0]
        with self.assertRaises(sqlite3.IntegrityError):self.command('responsibility.transfer.confirm',{'proposal_id':proposal['id']},1,
            evidence={'type':'external_reference','id':'synthetic:transfer'})
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_reasignaciones_obligacion').fetchone()[0],0)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_hechos_economicos').fetchone()[0],events)
        for rid in receipts:self.assertEqual(responsibility_balance(self.conn,self.community,rid)[0]['pending_cents'],'10000')

    def test_partial_many_collections_and_overpayment(self):
        rid=self.receipt();c1=self.collection(3000);self.allocate(c1,rid,3000)
        self.assertEqual(receipt_balance(self.conn,self.community,rid)['pending_cents'],'7000')
        c2=self.collection(9000,'BANK-2');self.allocate(c2,rid,7000)
        self.assertEqual(receipt_balance(self.conn,self.community,rid)['pending_cents'],'0')
        self.assertEqual(collection_balance(self.conn,self.community,c2)['available_cents'],'2000')
        self.assertEqual(receipt_balance(self.conn,self.community,rid,'2026-01-14')['pending_cents'],'10000')
        self.assertEqual(receipt_balance(self.conn,self.community,rid,'2026-01-20','2026-02-01T00:00:00Z')['pending_cents'],'10000')

    def test_multiple_receipts_and_stale_preview(self):
        a=self.receipt();b=self.receipt(15000,'TEST-2');cid=self.collection(20000)
        p=self.command('allocation.preview',{'collection_id':cid,'effective_on':'2026-01-15','allocations':[
            {'receipt_id':a,'amount_cents':'10000'},{'receipt_id':b,'amount_cents':'10000'}]})['entity']
        self.allocate(cid,a,1000)
        with self.assertRaises(ConflictError):self.command('allocation.confirm',{'proposal_id':p['id']},1)
        self.assertEqual(collection_balance(self.conn,self.community,cid)['available_cents'],'19000')
        self.assertEqual(receipt_balance(self.conn,self.community,b)['pending_cents'],'15000')

    def test_immutable_and_cross_community_fk(self):
        rid=self.receipt()
        with self.assertRaises(sqlite3.IntegrityError):self.conn.execute('UPDATE erp_recibos SET amount_cents=1 WHERE id=?',(rid,))
        with self.assertRaises(sqlite3.IntegrityError):self.conn.execute('DELETE FROM erp_recibos WHERE id=?',(rid,))
        cid=self.collection()
        with self.assertRaises(sqlite3.IntegrityError):self.conn.execute('DELETE FROM erp_cobros WHERE id=?',(cid,))
        other=self.conn.execute('SELECT id_comunidad FROM comunidades WHERE id_comunidad<>? LIMIT 1',(self.community,)).fetchone()[0]
        event=self.conn.execute('SELECT id FROM erp_hechos_economicos WHERE id_comunidad=? LIMIT 1',(self.community,)).fetchone()[0]
        with self.assertRaises(sqlite3.IntegrityError):
            self.conn.execute('''INSERT INTO erp_imputaciones (id_comunidad,collection_id,receipt_id,event_id,amount_cents,effective_on,registered_at)
                VALUES (?,?,?,?,100,'2026-01-15','2026-01-15T00:00:00.000000Z')''',(other,cid,rid,event))
        with self.assertRaises(sqlite3.IntegrityError):self.conn.execute('UPDATE erp_recibos SET id=id+1000000 WHERE id=?',(rid,))

    def test_read_only_user_and_no_implicit_economic_grants(self):
        row=self.conn.execute("SELECT id_usuario FROM usuarios WHERE activo=1 AND rol<>'Superusuario' LIMIT 1").fetchone()
        self.assertIsNotNone(row)
        user=profile(self.conn,row[0])
        p={'amount_cents':'100','currency':'EUR','effective_on':'2026-01-15','method':'efectivo','external_source':'test','external_key':'denied'}
        with self.assertRaises(PermissionError):self.command('collection.record',p,session=user)
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM erp_cobros').fetchone()[0],0)

    def test_idempotency_and_permissions(self):
        p={'amount_cents':'10000','currency':'EUR','effective_on':'2026-01-15','method':'efectivo','external_source':'test','external_key':'one'}
        a=self.command('collection.record',p,key='same');b=self.command('collection.record',p,key='same')
        self.assertEqual(a['entity']['id'],b['entity']['id'])
        self.assertTrue(b['idempotent_replay'])
        with self.assertRaises(ConflictError):self.command('collection.record',{**p,'amount_cents':'20000'},key='same')
        stale={**self.session,'auth_version':-1}
        with self.assertRaises(PermissionError):self.command('collection.record',{**p,'external_key':'other'},session=stale)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_cobros').fetchone()[0],1)

    def test_partial_return_keeps_original(self):
        rid=self.receipt();cid=self.collection();self.allocate(cid,rid,10000)
        aid=self.conn.execute('SELECT id FROM erp_imputaciones WHERE receipt_id=?',(rid,)).fetchone()[0]
        p=self.command('return.preview',{'collection_id':cid,'effective_on':'2026-01-20','free_cents':'0',
            'reversals':[{'allocation_id':aid,'amount_cents':'4000'}],'external_key':'RETURN-1'})['entity']
        self.command('return.confirm',{'proposal_id':p['id']},1)
        self.assertEqual(receipt_balance(self.conn,self.community,rid)['pending_cents'],'4000')
        self.assertEqual(collection_balance(self.conn,self.community,cid)['returned_cents'],'4000')
        self.assertEqual(self.conn.execute('SELECT amount_cents FROM erp_cobros WHERE id=?',(cid,)).fetchone()[0],10000)
        with self.assertRaises(ConflictError):self.command('void.preview',{'receipt_id':rid,'effective_on':'2026-01-21'})

    def test_credit_void_and_atomic_failure(self):
        rid=self.receipt()
        p=self.command('credit.preview',{'receipt_id':rid,'effective_on':'2026-01-10','amount_cents':'2000'})['entity']
        self.command('credit.confirm',{'proposal_id':p['id']},1)
        self.assertEqual(receipt_balance(self.conn,self.community,rid)['pending_cents'],'8000')
        second=self.receipt(10000,'VOID')
        p=self.command('void.preview',{'receipt_id':second,'effective_on':'2026-01-10'})['entity']
        self.command('void.confirm',{'proposal_id':p['id']},1)
        self.assertEqual(receipt_balance(self.conn,self.community,second)['state'],'anulado')
        third=self.receipt(10000,'ROLLBACK');cid=self.collection()
        p=self.command('allocation.preview',{'collection_id':cid,'effective_on':'2026-01-15','allocations':[{'receipt_id':third,'amount_cents':'1000'}]})['entity']
        before=self.conn.execute('SELECT count(*) FROM erp_hechos_economicos').fetchone()[0]
        self.conn.execute("CREATE TRIGGER injected_failure BEFORE INSERT ON erp_imputaciones BEGIN SELECT RAISE(ABORT,'test failure'); END")
        with self.assertRaises(sqlite3.IntegrityError):self.command('allocation.confirm',{'proposal_id':p['id']},1)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_hechos_economicos').fetchone()[0],before)
        self.assertEqual(receipt_balance(self.conn,self.community,third)['pending_cents'],'10000')


    def test_tabular_mapping_revision_exact_values_and_freeze(self):
        from erp_core.receivables_tabular import money_text
        self.assertEqual(money_text('1.234,56'),'123456')
        self.assertEqual(money_text('-0,01'),'-1')
        for value in ('1.234','1e3','1,234','12.34,56',float('nan')):
            with self.assertRaises(ContractError):money_text(value)
        code=self.conn.execute('SELECT codigo_propiedad FROM cf_propiedades WHERE id_propiedad=?',(self.property,)).fetchone()[0]
        payload={'source':'tabular-test','file_hash':'e'*64,'file_name':'historico.xlsx',
                 'rows':[{'rowNumber':2,'values':{'Referencia':'SALDO-1','Propiedad':code,'Saldo':'53,00'}}],
                 'mapping':{'Referencia':'reference','Propiedad':'property_code','Saldo':'amount'},
                 'defaults':{'cutoff_date':'2020-12-31','coverage_from':'2020-01-01','coverage_until':'2020-12-31',
                             'scope':'receipt','limitations':'Sin movimientos anteriores al corte'}}
        bad={**payload,'mapping':{**payload['mapping'],'Propiedad':''}}
        draft=self.command('history.import.preview',bad)['entity']
        self.assertEqual(draft['blocking_issues'],1)
        with self.assertRaises(ConflictError):self.command('history.import.preview',payload,999)
        draft=self.command('history.import.preview',payload,draft['version'])['entity']
        self.assertEqual(draft['blocking_issues'],0)
        self.assertEqual(draft['rows'][0]['normalized']['amount_cents'],'5300')
        self.assertEqual(draft['rows'][0]['normalized']['original']['row_number'],2)
        self.assertIsNotNone(draft['before'])
        self.command('history.import.confirm',{'import_id':draft['id']},draft['version'],evidence={'type':'external_reference','id':'test-file'})
        same=self.command('history.import.preview',payload)['entity']
        self.assertTrue(same['duplicate'])
        with self.assertRaises(ConflictError):self.command('history.import.preview',bad,same['version'])
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM erp_saldos_apertura').fetchone()[0],1)

    def test_workspace_and_scoped_collection_history(self):
        rid=self.receipt();cid=self.collection();self.allocate(cid,rid,3000)
        result=execute_query(self.db,self.session,{'query':'erp3.workspace.get','id_comunidad':self.community,'filters':{}})['entity']
        self.assertEqual(result['collection_count'],1)
        self.assertEqual(result['collections'][0]['balance']['available_cents'],'7000')
        detail=execute_query(self.db,self.session,{'query':'erp3.collection.get','id_comunidad':self.community,'filters':{'collection_id':cid}})['entity']
        self.assertEqual(detail['allocations'][0]['remaining_cents'],'3000')

    def opening(self,amount,reference='OPEN-1'):
        owner=self.conn.execute('SELECT id_propietario FROM cf_propietarios WHERE id_comunidad=? LIMIT 1',(self.community,)).fetchone()[0]
        row={'reference':reference,'amount_cents':str(amount),'property_id':self.property,'owner_id':owner,'attribution_confirmed':True,
             'cutoff_date':'2020-12-31','coverage_from':'2020-01-01','coverage_until':'2020-12-31','scope':'receipt','limitations':'Saldo documentado sin movimientos anteriores'}
        draft=self.command('history.import.preview',{'source':'opening-test','file_hash':hashlib.sha256(reference.encode()).hexdigest(),'file_name':'historico.csv','rows':[row]})['entity']
        result=self.command('history.import.confirm',{'import_id':draft['id']},draft['version'],evidence={'type':'external_reference','id':reference})['entity']
        return result['opening_ids'][0],owner

    def opening_move(self,payload):
        draft=self.command('opening.move.preview',payload)['entity']
        return self.command('opening.move.confirm',{'proposal_id':draft['id']},draft['version'],evidence={'type':'external_reference','id':'movimiento-documentado'})['entity']

    def test_opening_collection_and_reversal_preserve_source(self):
        from erp_core.receivables_projection import opening_balance
        oid,_=self.opening(530000);cid=self.collection(10000)
        before=tuple(self.conn.execute('SELECT * FROM erp_saldos_apertura WHERE id=?',(oid,)).fetchone())
        result=self.opening_move({'opening_id':oid,'kind':'allocation','effective_on':'2026-01-15','amount_cents':'3000','collection_id':cid})
        self.assertEqual(opening_balance(self.conn,self.community,oid)['remaining_cents'],'527000')
        self.assertEqual(collection_balance(self.conn,self.community,cid)['available_cents'],'7000')
        workspace=execute_query(self.db,self.session,{'query':'erp3.workspace.get','id_comunidad':self.community,'filters':{'property_id':self.property}})['entity']
        self.assertIn(cid,[r['id'] for r in workspace['collections']])
        self.assertEqual(opening_balance(self.conn,self.community,oid,'2025-01-01')['remaining_cents'],'530000')
        self.opening_move({'opening_id':oid,'kind':'reverse_allocation','effective_on':'2026-01-16','amount_cents':'1000','reverses_id':result['movement_id']})
        self.assertEqual(opening_balance(self.conn,self.community,oid)['remaining_cents'],'528000')
        self.assertEqual(collection_balance(self.conn,self.community,cid)['available_cents'],'8000')
        self.assertEqual(tuple(self.conn.execute('SELECT * FROM erp_saldos_apertura WHERE id=?',(oid,)).fetchone()),before)
        with self.assertRaises(ConflictError):self.command('opening.move.preview',{'opening_id':oid,'kind':'allocation','effective_on':'2026-01-17','amount_cents':'8001','collection_id':cid})

    def test_opening_credit_application_and_refund(self):
        from erp_core.receivables_projection import opening_balance
        oid,owner=self.opening(-10000);rid=self.receipt()
        move=self.opening_move({'opening_id':oid,'kind':'credit_apply','receipt_id':rid,'effective_on':'2026-01-16','amount_cents':'3000'})
        self.assertEqual(receipt_balance(self.conn,self.community,rid)['pending_cents'],'7000')
        self.assertEqual(opening_balance(self.conn,self.community,oid)['remaining_cents'],'-7000')
        self.opening_move({'opening_id':oid,'kind':'reverse_credit_apply','reverses_id':move['movement_id'],'effective_on':'2026-01-17','amount_cents':'1000'})
        self.opening_move({'opening_id':oid,'kind':'refund','effective_on':'2026-01-17','amount_cents':'2000','beneficiary':{'type':'owner','id':owner}})
        self.assertEqual(opening_balance(self.conn,self.community,oid)['remaining_cents'],'-6000')
        self.assertEqual(receipt_balance(self.conn,self.community,rid)['pending_cents'],'8000')
        self.assertEqual(self.conn.execute('SELECT COUNT(*) FROM erp_cobros').fetchone()[0],0)
        with self.assertRaises(ConflictError):self.command('void.preview',{'receipt_id':rid,'effective_on':'2026-01-18'})

    def test_return_reversal_restores_free_funds_without_autoallocation(self):
        rid=self.receipt();cid=self.collection();self.allocate(cid,rid,10000)
        aid=self.conn.execute('SELECT id FROM erp_imputaciones WHERE receipt_id=?',(rid,)).fetchone()[0]
        draft=self.command('return.preview',{'collection_id':cid,'effective_on':'2026-01-16','free_cents':'0','external_key':'return-then-correct','reversals':[{'allocation_id':aid,'amount_cents':'4000'}]})['entity']
        self.command('return.confirm',{'proposal_id':draft['id']},1)
        original=dict(self.conn.execute('SELECT * FROM erp_devoluciones WHERE collection_id=?',(cid,)).fetchone())
        draft=self.command('return.reverse.preview',{'return_id':original['id'],'effective_on':'2026-01-17'})['entity']
        self.command('return.reverse.confirm',{'proposal_id':draft['id']},1,evidence={'type':'external_reference','id':'banco-corrige-devolucion'})
        self.assertEqual(collection_balance(self.conn,self.community,cid,'2026-01-16')['returned_cents'],'4000')
        self.assertEqual(collection_balance(self.conn,self.community,cid)['returned_cents'],'0')
        self.assertEqual(collection_balance(self.conn,self.community,cid)['available_cents'],'4000')
        self.assertEqual(receipt_balance(self.conn,self.community,rid)['pending_cents'],'4000')
        self.assertEqual(dict(self.conn.execute('SELECT * FROM erp_devoluciones WHERE id=?',(original['id'],)).fetchone()),original)
        with self.assertRaises(ConflictError):self.command('return.reverse.preview',{'return_id':original['id'],'effective_on':'2026-01-18'})

    def test_unknown_payer_accreditation_is_temporal_and_preserves_original(self):
        from erp_core.receivables_projection import collection_payer
        cid=self.collection()
        owner=self.conn.execute('SELECT id_propietario FROM cf_propietarios WHERE id_comunidad=? LIMIT 1',(self.community,)).fetchone()[0]
        draft=self.command('collection.payer.preview',{'collection_id':cid,'payer':{'type':'owner','id':owner},'effective_on':'2026-01-16'})['entity']
        self.command('collection.payer.confirm',{'proposal_id':draft['id']},1,evidence={'type':'external_reference','id':'justificante-del-pagador'})
        self.assertIsNone(collection_payer(self.conn,self.community,cid,'2026-01-15')['payer_owner_id'])
        self.assertEqual(collection_payer(self.conn,self.community,cid,'2026-01-16')['payer_owner_id'],owner)
        self.assertIsNone(self.conn.execute('SELECT payer_owner_id FROM erp_cobros WHERE id=?',(cid,)).fetchone()[0])
        draft=self.command('refund.preview',{'collection_id':cid,'beneficiary':{'type':'owner','id':owner},'effective_on':'2026-01-17','amount_cents':'1000'})['entity']
        self.command('refund.confirm',{'proposal_id':draft['id']},1,evidence={'type':'external_reference','id':'justificante-reintegro'})
        self.assertEqual(collection_balance(self.conn,self.community,cid)['available_cents'],'9000')

    def test_credit_reversal_preserves_original_and_cutoffs(self):
        rid=self.receipt()
        draft=self.command('credit.preview',{'receipt_id':rid,'effective_on':'2026-01-16','amount_cents':'3000'})['entity']
        self.command('credit.confirm',{'proposal_id':draft['id']},1)
        original=dict(self.conn.execute("SELECT * FROM erp_rectificaciones WHERE receipt_id=? AND kind='credit'",(rid,)).fetchone())
        payload={'movement_id':original['id'],'effective_on':'2026-01-18','amount_cents':'1000'}
        draft=self.command('credit.reverse.preview',payload)['entity']
        with self.assertRaises(ContractError):self.command('credit.reverse.confirm',{'proposal_id':draft['id']},1)
        result=self.command('credit.reverse.confirm',{'proposal_id':draft['id']},1,key='reverse-credit-once',evidence={'type':'external_reference','id':'rectificacion-documentada'})
        replay=self.command('credit.reverse.confirm',{'proposal_id':draft['id']},1,key='reverse-credit-once',evidence={'type':'external_reference','id':'rectificacion-documentada'})
        self.assertEqual(result['entity'],replay['entity'])
        self.assertTrue(replay['idempotent_replay'])
        self.assertEqual(receipt_balance(self.conn,self.community,rid,'2026-01-17')['pending_cents'],'7000')
        self.assertEqual(receipt_balance(self.conn,self.community,rid,'2026-01-18')['pending_cents'],'8000')
        self.assertEqual(dict(self.conn.execute('SELECT * FROM erp_rectificaciones WHERE id=?',(original['id'],)).fetchone()),original)
        with self.assertRaises(ConflictError):self.command('credit.reverse.preview',{**payload,'amount_cents':'2001'})

    def test_credit_application_reversal_does_not_create_cash(self):
        from erp_core.receivables_projection import credit_balance
        rid=self.receipt()
        owner=self.conn.execute('SELECT id_propietario FROM cf_propietarios WHERE id_comunidad=? LIMIT 1',(self.community,)).fetchone()[0]
        event=self.conn.execute("INSERT INTO erp_hechos_economicos (id_comunidad,event_key,event_type,schema_version,effective_on,registered_at,actor_id,reason,payload_json,payload_hash,evidence_json) VALUES (?,'synthetic-credit','synthetic','erp_receivables_v1','2026-01-01','2026-01-01T00:00:00.000000Z',?,'test','{}','test','null')",(self.community,self.uid)).lastrowid
        credit=self.conn.execute("INSERT INTO erp_creditos (id_comunidad,event_id,owner_id,amount_cents,effective_on,registered_at) VALUES (?,?,?,4000,'2026-01-01','2026-01-01T00:00:00.000000Z')",(self.community,event,owner)).lastrowid
        draft=self.command('credit.apply.preview',{'credit_id':credit,'receipt_id':rid,'amount_cents':'3000','effective_on':'2026-01-16'})['entity']
        self.command('credit.apply.confirm',{'proposal_id':draft['id']},1)
        original=self.conn.execute('SELECT id FROM erp_credito_aplicaciones WHERE credit_id=? AND reverses_id IS NULL',(credit,)).fetchone()[0]
        draft=self.command('credit.apply.reverse.preview',{'movement_id':original,'amount_cents':'1000','effective_on':'2026-01-17'})['entity']
        self.command('credit.apply.reverse.confirm',{'proposal_id':draft['id']},1,evidence={'type':'external_reference','id':'reversion-documentada'})
        self.assertEqual(credit_balance(self.conn,self.community,credit)['available_cents'],'2000')
        self.assertEqual(receipt_balance(self.conn,self.community,rid)['pending_cents'],'8000')
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_cobros').fetchone()[0],0)

    def test_period_summary_uses_cash_once_and_respects_cuts(self):
        rid=self.receipt();cid=self.collection(12000);self.allocate(cid,rid,10000)
        aid=self.conn.execute('SELECT id FROM erp_imputaciones WHERE receipt_id=?',(rid,)).fetchone()[0]
        draft=self.command('return.preview',{'collection_id':cid,'effective_on':'2026-01-20','free_cents':'1000','external_key':'mixed-return',
            'reversals':[{'allocation_id':aid,'amount_cents':'3000'}]})['entity']
        self.command('return.confirm',{'proposal_id':draft['id']},1)
        report=execute_query(self.db,self.session,{'query':'erp3.period.summary','id_comunidad':self.community,
            'filters':{'from':'2026-01-01','until':'2026-01-31'}})['entity']
        self.assertEqual(report['issued_cents'],'10000')
        self.assertEqual(report['cash_received_cents'],'12000')
        self.assertEqual(report['cash_returns_cents'],'4000')
        self.assertEqual(report['issued_in_period_pending_at_end_cents'],'3000')
        self.assertFalse(report['legacy_combined'])
        before=execute_query(self.db,self.session,{'query':'erp3.period.summary','id_comunidad':self.community,
            'filters':{'from':'2026-01-01','until':'2026-01-19'}})['entity']
        self.assertEqual(before['cash_returns_cents'],'0')
        self.assertEqual(before['issued_in_period_pending_at_end_cents'],'0')

    def test_statement_and_assistant_share_exact_projection(self):
        from erp_core.receivables_agent import debt_answer
        oid,owner=self.opening(5300)
        rid=self.receipt();cid=self.collection();self.allocate(cid,rid,3000)
        statement=execute_query(self.db,self.session,{'query':'erp3.account.statement','id_comunidad':self.community,
            'filters':{'property_id':self.property}})['entity']
        native=next(r for r in statement['items'] if r.get('receipt_id')==rid)
        self.assertEqual(native['pending_cents'],'7000')
        opening=next(r for r in statement['items'] if r.get('opening_id')==oid)
        self.assertEqual(opening['pending_cents'],'5300')
        answer=debt_answer(self.db,self.session,[self.community],property_id=self.property)
        self.assertEqual(answer['facts']['deuda_centimos'],statement['documented_subtotal_cents'])
        self.assertTrue(any(r['Referencia']=='TEST-1' and r['Pendiente']=='70,00 EUR' for r in answer['display']['tables'][0]['rows']))
        personal=debt_answer(self.db,self.session,[self.community],owner_id=owner)
        self.assertNotIn('no tiene deuda',personal['answer'])
        self.assertEqual(personal['data_status'],'incompleto')

if __name__=='__main__':
    print('Isolated workspace:',WORK,flush=True)
    unittest.main(verbosity=2)
