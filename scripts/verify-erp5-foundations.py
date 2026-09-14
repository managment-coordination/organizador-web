"""ERP 5 acceptance on independent copies and public synthetic bank fixtures."""
import base64
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from datetime import datetime, timezone
import importlib.util
import json
from pathlib import Path
import sqlite3
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'server'))
from access_control import profile
from erp_core.banking_crypto import BankVault
from erp_core.banking_service import BankingService
from erp_core.contracts import CommandEnvelope, QueryEnvelope
from erp_core.database import connect
from erp_core.errors import ContractError,ConflictError
from erp_core.migrations import apply_all,MIGRATIONS
from erp_core.receivables_service import ReceivablesService
from erp_core.reconciliation_service import ReconciliationService,CAPABILITIES
from erp_core.reconciliation_adapters import amount,parse

SOURCE=Path(sys.argv[1]).resolve();sys.argv=[sys.argv[0],*sys.argv[2:]]
WORK=Path(tempfile.mkdtemp(prefix='organizador-erp5-foundations-'))


class ReconciliationTests(unittest.TestCase):
    def setUp(self):
        self.work=WORK/self._testMethodName;self.work.mkdir()
        self.db=self.work/'database.db'
        with closing(sqlite3.connect('file:'+str(SOURCE)+'?mode=ro',uri=True)) as src,closing(sqlite3.connect(self.db)) as dst:src.backup(dst)
        self.conn=connect(self.db);apply_all(self.conn)
        self.uid=self.conn.execute("SELECT id_usuario FROM usuarios WHERE rol='Superusuario' AND activo=1 LIMIT 1").fetchone()[0]
        self.session=profile(self.conn,self.uid);self.session['banking_reauthenticated_at']=datetime.now(timezone.utc).isoformat()
        self.community=self.session['comunidades'][0]['id_comunidad'];self.seq=0
        self.vault=BankVault(BankVault.create_key_file(self.work/'custody.key'))
        self.service=ReconciliationService(self.db,vault=self.vault)
        for cap in CAPABILITIES:self.service.permissions_save(self.session,self.e('permissions.save',{'user_id':self.uid,'capability':cap,'allowed':True},0))
        bank=BankingService(self.db,vault=self.vault)
        bank.permissions_save(self.session,self.e4('permissions.save',{'user_id':self.uid,'capability':'configure_creditor','allowed':True},0))
        code='ES'+str(98-int('M12345678'.replace('M','22')+'142800')%97).zfill(2)+'000M12345678'
        self.account=bank.creditor_create(self.session,self.e4('creditor.create',{'name':'Comunidad sintetica','creditor_identifier':code,
            'iban':'ES9121000418450200051332','address':{'country':'ES','town':'Madrid'},'effective_from':'2026-01-01'}))['entity']['treasury_id']

    def tearDown(self):
        self.assertEqual(self.conn.execute('PRAGMA integrity_check').fetchone()[0],'ok')
        self.assertEqual(self.conn.execute('PRAGMA foreign_key_check').fetchall(),[])
        self.conn.close()

    def e(self,name,p,version=None,key=None):
        self.seq+=1
        return CommandEnvelope.from_value({'command':'erp5.'+name,'id_comunidad':self.community,'payload':p,
            'idempotency_key':key or 'synthetic-'+str(self.seq),'expected_version':version,'origin':'test','reason':'Evidencia sintetica',
            'evidence':{'type':'external_reference','id':'synthetic-evidence'}})

    def e4(self,name,p,version=None):
        from dataclasses import replace
        return replace(self.e(name,p,version),command='erp4.'+name)

    def q(self,name,p):return QueryEnvelope('erp5.'+name,self.community,p)

    def imported(self,rows,coverage=None):
        p={'treasury_id':self.account,'profile':'manual-v1','rows':rows}
        if coverage:p['coverage']=coverage
        proposal=self.service.statement_preview(self.session,self.e('statement.preview',p))['entity']
        confirmed=self.service.statement_confirm(self.session,self.e('statement.confirm',{'id':proposal['id']},proposal['version']))['entity']
        return confirmed['movement_ids']

    def row(self,n='10000',ref='ID-1',on='2026-09-01'):
        return {'operation_on':on,'value_on':'2026-09-02','amount_cents':n,'currency':'EUR','concept':'Prueba sintetica',
                'external_id':ref,'namespace':'synthetic-bank'}

    def collection(self,n='10000'):
        e=self.e('collection.record',{'amount_cents':n,'currency':'EUR','effective_on':'2026-09-01','method':'transferencia',
            'external_source':'synthetic','external_key':'collection-'+str(self.seq),'treasury_reference':'erp4-treasury:'+str(self.account)})
        from dataclasses import replace
        return ReceivablesService(self.db).collection_record(self.session,replace(e,command='erp3.collection.record'))['entity']['id']

    def match(self,components,net=None):
        p={'components':components}
        if net:p['net_evidence']=net
        preview=self.service.match_preview(self.session,self.e('match.preview',p))['entity']
        return self.service.match_confirm(self.session,self.e('match.confirm',{'proposal_id':preview['id']},preview['version']))['entity']

    def test_01_migration_reentrant_history(self):
        checks=list(self.conn.execute('SELECT version,checksum FROM erp_schema_migrations'));apply_all(self.conn)
        self.assertEqual(checks,list(self.conn.execute('SELECT version,checksum FROM erp_schema_migrations')))
        self.assertEqual(checks[-1][0],20)

    def test_02_duplicate_file_and_stable_id(self):
        first=self.imported([self.row()]);again=self.imported([self.row()])
        self.assertEqual(first,again)
        second=self.imported([self.row(),self.row('20000','ID-2')]);self.assertEqual(second[0],first[0])
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_banco_movimientos').fetchone()[0],2)

    def test_03_equal_legitimate_movements(self):
        r=self.row(ref='');ids=self.imported([r,r]);self.assertEqual(len(set(ids)),2)
        preview=self.service.statement_preview(self.session,self.e('statement.preview',{'treasury_id':self.account,'rows':[r], 'profile':'manual-v1'}))['entity']
        self.assertEqual(preview['review_count'],1)
        with self.assertRaises(ConflictError):self.service.statement_confirm(self.session,self.e('statement.confirm',{'id':preview['id']},1))

    def test_04_conflicting_identity(self):
        self.imported([self.row()])
        with self.assertRaises(ConflictError):self.imported([self.row('12000')])

    def test_05_link_existing_collection_no_funds(self):
        mid=self.imported([self.row()])[0];cid=self.collection()
        before=self.conn.execute('SELECT count(*) FROM erp_cobros').fetchone()[0]
        self.match([{'movement_id':mid,'action':'link_collection','fact_id':cid,'amount_cents':'10000'}])
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_cobros').fetchone()[0],before)
        self.assertEqual(self.service.movement_get(self.session,self.q('movement.get',{'id':mid}))['entity']['state'],'conciliado')

    def test_06_partial_assignment(self):
        mid=self.imported([self.row()])[0]
        self.match([{'movement_id':mid,'action':'record_collection','amount_cents':'6000'}])
        self.assertEqual(self.service.movement_get(self.session,self.q('movement.get',{'id':mid}))['entity']['remaining_cents'],'4000')

    def test_07_many_movements_one_fact(self):
        mids=self.imported([self.row('3000','ID-1'),self.row('7000','ID-2')]);cid=self.collection()
        self.match([{'movement_id':mid,'action':'link_collection','fact_id':cid,'amount_cents':str(n)} for mid,n in zip(mids,[3000,7000])])
        with self.assertRaises(ConflictError):self.match([{'movement_id':mids[0],'action':'link_collection','fact_id':cid,'amount_cents':'1'}])

    def test_08_net_commission_exact(self):
        mid=self.imported([self.row('98000')])[0]
        self.match([{'movement_id':mid,'action':'record_collection','amount_cents':'100000'},
            {'movement_id':mid,'action':'record_outflow','amount_cents':'-2000','kind':'commission','description':'Comision acreditada'}],net='synthetic-net-document')
        self.assertEqual(self.service.movement_get(self.session,self.q('movement.get',{'id':mid}))['entity']['remaining_cents'],'0')

    def test_09_no_silent_net(self):
        mid=self.imported([self.row('98000')])[0]
        with self.assertRaises(ConflictError):self.match([{'movement_id':mid,'action':'record_collection','amount_cents':'100000'}])

    def test_10_reverse_preserves_economic_history(self):
        mid=self.imported([self.row()])[0];match=self.match([{'movement_id':mid,'action':'record_collection','amount_cents':'10000'}])
        before=self.conn.execute('SELECT count(*) FROM erp_cobros').fetchone()[0]
        self.service.match_reverse(self.session,self.e('match.reverse',{'id':match['id']}))
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_cobros').fetchone()[0],before)
        self.assertEqual(self.service.movement_get(self.session,self.q('movement.get',{'id':mid}))['entity']['remaining_cents'],'10000')

    def test_11_unknown_and_operation_value(self):
        mid=self.imported([self.row()])[0];r=self.service.movement_get(self.session,self.q('movement.get',{'id':mid}))['entity']
        self.assertEqual(r['state'],'pendiente');self.assertNotEqual(r['operation_on'],r['value_on'])

    def test_12_exact_decimal_rejection(self):
        self.assertEqual(amount('1.234,56',','),'123456')
        with self.assertRaises(ContractError):amount('0.001')
        with self.assertRaises(ContractError):amount(1.01)

    def test_13_permissions_no_superuser_implicit(self):
        self.service.permissions_save(self.session,self.e('permissions.save',{'user_id':self.uid,'capability':'import','allowed':False},1))
        with self.assertRaises(PermissionError):self.imported([self.row()])

    def test_14_idempotence(self):
        mid=self.imported([self.row()])[0]
        preview=self.service.match_preview(self.session,self.e('match.preview',{'components':[{'movement_id':mid,'action':'record_collection','amount_cents':'10000'}]}))['entity']
        e=self.e('match.confirm',{'proposal_id':preview['id']},1,key='repeat')
        one=self.service.match_confirm(self.session,e);two=self.service.match_confirm(self.session,e)
        self.assertEqual(one['entity'],two['entity']);self.assertTrue(two['idempotent_replay'])

    def test_15_cannot_double_consume(self):
        mid=self.imported([self.row()])[0]
        self.match([{'movement_id':mid,'action':'record_collection','amount_cents':'10000'}])
        with self.assertRaises(ConflictError):self.match([{'movement_id':mid,'action':'record_collection','amount_cents':'1'}])

    def test_16_closure_pending_reopen(self):
        coverage={'start_on':'2026-09-01','end_on':'2026-09-30','opening_cents':'1000','closing_cents':'11000','complete':True,'balance_type':'booked'}
        mid=self.imported([self.row()],coverage)[0]
        self.service.pending_assign(self.session,self.e('pending.assign',{'movement_id':mid,'user_id':self.uid,'review_on':'2026-10-01','note':'Falta identificar pagador'}))
        p=self.service.closure_preview(self.session,self.e('closure.preview',{'treasury_id':self.account,'start_on':'2026-09-01','end_on':'2026-09-30'}))['entity']
        c=self.service.closure_confirm(self.session,self.e('closure.confirm',{'proposal_id':p['id']},1))['entity']
        self.assertEqual(c['state'],'cerrado_pendientes')
        with self.assertRaises(ConflictError):self.match([{'movement_id':mid,'action':'record_collection','amount_cents':'10000'}])
        self.service.closure_reopen(self.session,self.e('closure.reopen',{'id':c['id']},1))
        self.match([{'movement_id':mid,'action':'record_collection','amount_cents':'10000'}])

    def test_17_difference_blocks_closure(self):
        self.imported([self.row()],{'start_on':'2026-09-01','end_on':'2026-09-30','opening_cents':'0','closing_cents':'9999','complete':True})
        with self.assertRaises(ConflictError):self.service.closure_preview(self.session,self.e('closure.preview',{'treasury_id':self.account,'start_on':'2026-09-01','end_on':'2026-09-30'}))

    def test_18_pending_source_no_funds(self):
        mid=self.imported([{**self.row(),'source_state':'pending'}])[0]
        with self.assertRaises(ContractError):self.match([{'movement_id':mid,'action':'record_collection','amount_cents':'10000'}])

    def test_19_ciphertext_logs_and_events(self):
        r={**self.row(),'concept':'IBAN ES9121000418450200051332'};mid=self.imported([r])[0]
        public=self.service.movement_get(self.session,self.q('movement.get',{'id':mid}))['entity']
        self.assertNotIn('ES9121000418450200051332',json.dumps(public))
        for table in ('erp_banca_secretos','erp_audit_events','erp_outbox','erp_command_log'):
            self.assertNotIn('ES9121000418450200051332',repr([tuple(r) for r in self.conn.execute('SELECT * FROM '+table)]))

    def test_20_payment_not_contability(self):
        mid=self.imported([self.row('-10000')])[0]
        self.match([{'movement_id':mid,'action':'record_outflow','amount_cents':'-10000','kind':'payment','description':'Pago sintetico'}])
        events=list(self.conn.execute("SELECT payload_json FROM erp_outbox WHERE event_type='erp5.outflow.confirmed'"))
        self.assertEqual(len(events),1);self.assertIn('monetary',events[0][0])

    def second_account(self,kind='banco'):
        secret=self.vault.put(self.conn,self.community,'treasury-account',{'iban':'ES6621000418401234567891'},'2026-09-01T00:00:00Z') if kind=='banco' else None
        with self.conn:
            return self.conn.execute('''INSERT INTO erp_cuentas_tesoreria (id_comunidad,kind,currency,state,secret_id,registered_at,actor_id)
                VALUES (?,?,'EUR','activa',?,?,?)''',(self.community,kind,secret,'2026-09-01T00:00:00Z',self.uid)).lastrowid

    def transfer(self,p):
        preview=self.service.transfer_preview(self.session,self.e('transfer.preview',p))['entity']
        return self.service.transfer_confirm(self.session,self.e('transfer.confirm',{'proposal_id':preview['id']},preview['version']))['entity']

    def test_21_transfer_both_sides(self):
        other=self.second_account();source=self.imported([self.row('-10000')])[0]
        account=self.account;self.account=other;destination=self.imported([self.row('10000','ID-2')])[0];self.account=account
        before=self.conn.execute('SELECT count(*) FROM erp_cobros').fetchone()[0]
        t=self.transfer({'source_id':account,'destination_id':other,'amount_cents':'10000','legs':[
            {'side':'source','movement_id':source,'amount_cents':'10000'},{'side':'destination','movement_id':destination,'amount_cents':'10000'}]})
        self.assertEqual(t['state'],'confirmada');self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_cobros').fetchone()[0],before)
        events=[json.loads(r[0]) for r in self.conn.execute("SELECT payload_json FROM erp_outbox WHERE event_type='erp5.transfer.leg_confirmed'")]
        self.assertEqual(sum(int(r['amount_cents']) for r in events),0);self.assertEqual(len(events),2)

    def test_22_transfer_in_transit_and_completion(self):
        other=self.second_account();mid=self.imported([self.row('-10000')])[0]
        t=self.transfer({'source_id':self.account,'destination_id':other,'amount_cents':'10000','legs':[{'side':'source','movement_id':mid,'amount_cents':'10000'}]})
        self.assertEqual(t['state'],'en_transito')
        account=self.account;self.account=other;destination=self.imported([self.row('10000','ID-2')])[0];self.account=account
        p={'transfer_id':t['id'],'legs':[{'side':'destination','movement_id':destination,'amount_cents':'10000'}]}
        preview=self.service.transfer_complete_preview(self.session,self.e('transfer.complete_preview',p))['entity']
        result=self.service.transfer_complete_confirm(self.session,self.e('transfer.complete_confirm',{'proposal_id':preview['id']},preview['version']))['entity']
        self.assertEqual(result['state'],'confirmada')
        self.assertEqual(self.service.transfer_list(self.session,self.q('transfer.list',{}))['entity']['items'][0]['state'],'confirmada')
        with self.assertRaises((ConflictError,ContractError)):self.service.transfer_complete_preview(self.session,self.e('transfer.complete_preview',p))

    def test_23_cash_transfer_requires_evidence(self):
        other=self.second_account('caja');mid=self.imported([self.row('-10000')])[0]
        p={'source_id':self.account,'destination_id':other,'amount_cents':'10000','legs':[
            {'side':'source','movement_id':mid,'amount_cents':'10000'},
            {'side':'destination','amount_cents':'10000','effective_on':'2026-09-01','cash_evidence':'Justificante sintetico'}]}
        t=self.transfer(p);self.assertEqual(t['state'],'confirmada')
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_banco_movimientos').fetchone()[0],1)

    def test_24_rollback_after_economic_creation(self):
        mid=self.imported([self.row()])[0]
        preview=self.service.match_preview(self.session,self.e('match.preview',{'components':[{'movement_id':mid,'action':'record_collection','amount_cents':'10000'}]}))['entity']
        baseline={t:self.conn.execute('SELECT count(*) FROM '+t).fetchone()[0] for t in ('erp_cobros','erp_conciliaciones','erp_conciliacion_componentes','erp_audit_events','erp_outbox')}
        with patch.object(self.service,'_event',side_effect=RuntimeError('Synthetic failure')):
            with self.assertRaises(RuntimeError):self.service.match_confirm(self.session,self.e('match.confirm',{'proposal_id':preview['id']},1))
        self.assertEqual(baseline,{t:self.conn.execute('SELECT count(*) FROM '+t).fetchone()[0] for t in baseline})

    def test_25_pending_to_booked_preserves_history(self):
        old=self.imported([{**self.row(),'source_state':'pending'}])[0]
        new=self.imported([self.row()])[0];self.assertNotEqual(old,new)
        self.assertEqual(self.service.movement_get(self.session,self.q('movement.get',{'id':old}))['entity']['state'],'rectificado')
        self.assertEqual(self.imported([self.row()]),[new])
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_banco_movimientos').fetchone()[0],2)

    def test_26_return_uses_erp3_and_reversal_does_not_restore_funds(self):
        cid=self.collection();mid=self.imported([self.row('-4000')])[0]
        match=self.match([{'movement_id':mid,'action':'return','amount_cents':'-4000','return_payload':{
            'collection_id':cid,'effective_on':'2026-09-01','free_cents':'4000','reversals':[],'external_key':'synthetic-return'}}])
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_devoluciones').fetchone()[0],1)
        self.service.match_reverse(self.session,self.e('match.reverse',{'id':match['id']}))
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_devoluciones').fetchone()[0],1)

    def test_27_concurrent_reviewers(self):
        mid=self.imported([self.row()])[0]
        previews=[self.service.match_preview(self.session,self.e('match.preview',{'components':[{'movement_id':mid,'action':'record_collection','amount_cents':'10000'}]}))['entity'] for _ in range(2)]
        envelopes=[self.e('match.confirm',{'proposal_id':p['id']},1) for p in previews]
        def confirm(e):
            try:self.service.match_confirm(self.session,e);return 'confirmed'
            except ConflictError:return 'conflict'
        with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(confirm,envelopes))
        self.assertEqual(sorted(results),['confirmed','conflict'])

    def test_28_structural_history_and_exact_amount(self):
        mid=self.imported([self.row()])[0]
        with self.assertRaises(sqlite3.IntegrityError):self.conn.execute('UPDATE erp_banco_movimientos SET amount_cents=1 WHERE id=?',(mid,))
        self.conn.rollback()
        with self.assertRaises(sqlite3.IntegrityError):self.conn.execute('DELETE FROM erp_banco_movimientos WHERE id=?',(mid,))
        self.conn.rollback()

    def test_29_activation_routes_without_legacy_double_sum(self):
        from erp_core.reconciliation_sources import period_bank_answer
        coverage={'start_on':'2026-09-01','end_on':'2026-09-30','opening_cents':'1000','closing_cents':'11000','complete':True,'balance_type':'booked'}
        mid=self.imported([self.row()],coverage)[0]
        self.service.pending_assign(self.session,self.e('pending.assign',{'movement_id':mid,'user_id':self.uid,'review_on':'2026-10-01','note':'Pendiente documentado'}))
        p={'treasury_id':self.account,'start_on':'2026-09-01','end_on':'2026-09-30'}
        self.assertIsNone(period_bank_answer(self.db,self.session,[self.community],p['start_on'],p['end_on']))
        preview=self.service.source_preview(self.session,self.e('source.preview',p))['entity']
        self.service.source_confirm(self.session,self.e('source.confirm',{'proposal_id':preview['id']},1))
        answer=period_bank_answer(self.db,self.session,[self.community],p['start_on'],p['end_on'])
        self.assertEqual(answer['display_end'],'110,00 EUR');self.assertEqual(answer['source'],'ERP5')
        with self.assertRaises(ConflictError):self.service.source_preview(self.session,self.e('source.preview',p))
        self.service.permissions_save(self.session,self.e('permissions.save',{'user_id':self.uid,'capability':'read','allowed':False},1))
        answer=period_bank_answer(self.db,self.session,[self.community],p['start_on'],p['end_on'])
        self.assertFalse(answer['complete']);self.assertEqual(answer['display_end'],'No verificable')

    def test_30_export_masked_exact_audited_and_formula_safe(self):
        import io
        from openpyxl import load_workbook
        self.imported([{**self.row(),'concept':'=HYPERLINK("test") ES9121000418450200051332'}])
        p={'treasury_id':self.account,'start_on':'2026-09-01','end_on':'2026-09-30','format':'xlsx'}
        report=self.service.report_export(self.session,self.e('report.export',p))['entity']
        book=load_workbook(io.BytesIO(base64.b64decode(report['content_base64'])),data_only=False)
        self.assertEqual(book.active['D2'].value,'100.00');self.assertEqual(book.active['C2'].data_type,'s')
        self.assertNotIn('ES9121000418450200051332',book.active['C2'].value);book.close()
        self.assertEqual(self.conn.execute("SELECT count(*) FROM erp_outbox WHERE event_type='erp5.report.exported'").fetchone()[0],1)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_cobros').fetchone()[0],0)

    def test_31_closure_blocks_spanning_coverage(self):
        p={'treasury_id':self.account,'start_on':'2026-09-10','end_on':'2026-09-20','opening_cents':'0','closing_cents':'0','complete':True}
        self.service.opening_confirm(self.session,self.e('opening.confirm',p))
        preview=self.service.closure_preview(self.session,self.e('closure.preview',{k:p[k] for k in ('treasury_id','start_on','end_on')}))['entity']
        self.service.closure_confirm(self.session,self.e('closure.confirm',{'proposal_id':preview['id']},1))
        with self.assertRaises(ConflictError):self.service.opening_confirm(self.session,self.e('opening.confirm',{**p,'start_on':'2026-09-01','end_on':'2026-09-30'}))

    def test_32_link_return_not_second_debt(self):
        cid=self.collection();mid=self.imported([self.row('-4000')])[0]
        result=self.match([{'movement_id':mid,'action':'return','amount_cents':'-4000','return_payload':{
            'collection_id':cid,'effective_on':'2026-09-01','free_cents':'4000','reversals':[],'external_key':'RETURN-ONE'}}])
        returned=self.conn.execute('SELECT id FROM erp_devoluciones').fetchone()[0]
        self.service.match_reverse(self.session,self.e('match.reverse',{'id':result['id']}))
        self.match([{'movement_id':mid,'action':'link_return','amount_cents':'-4000','fact_id':returned}])
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_devoluciones').fetchone()[0],1)

    def test_33_same_stable_id_twice_in_file_is_not_two_movements(self):
        with self.assertRaises(ConflictError):self.service.statement_preview(self.session,self.e('statement.preview',{
            'treasury_id':self.account,'profile':'manual-v1','rows':[self.row(),self.row()]}))
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_banco_movimientos').fetchone()[0],0)

    def test_34_original_requires_reauth_and_separate_bank_grants(self):
        raw=b'Fecha;Importe;Concepto\n2026-09-01;100.00;Sintetico\n'
        p=self.service.statement_preview(self.session,self.e('statement.preview',{'treasury_id':self.account,'profile':'tabular-v1',
            'filename':'original.csv','content_base64':base64.b64encode(raw).decode(),'mapping':{'operation_on':0,'amount':1,'concept':2},
            'options':{'decimal_separator':'.','date_format':'iso'}}))['entity']
        with self.assertRaises(PermissionError):self.service.statement_original(self.session,self.community,p['id'],'Revision sintetica')
        bank=BankingService(self.db,vault=self.vault)
        for cap in ('export','reveal'):bank.permissions_save(self.session,self.e4('permissions.save',{'user_id':self.uid,'capability':cap,'allowed':True},0))
        stale={**self.session,'banking_reauthenticated_at':'2000-01-01T00:00:00+00:00'}
        with self.assertRaises(PermissionError):self.service.statement_original(stale,self.community,p['id'],'Revision sintetica')
        result=self.service.statement_original(self.session,self.community,p['id'],'Revision sintetica')
        self.assertEqual(base64.b64decode(result['content_base64']),raw)
        self.assertEqual(self.conn.execute("SELECT count(*) FROM erp_audit_events WHERE action='erp5.statement.original'").fetchone()[0],1)
        audit=' '.join(str(tuple(r)) for r in self.conn.execute('SELECT * FROM erp_audit_events'))
        self.assertNotIn(base64.b64encode(raw).decode(),audit)

    def test_35_refund_delegates_once_and_keeps_free_balance(self):
        from dataclasses import replace
        owner=self.conn.execute('SELECT id_propietario FROM cf_propietarios WHERE id_comunidad=? LIMIT 1',(self.community,)).fetchone()[0]
        payload={'amount_cents':'10000','currency':'EUR','effective_on':'2026-09-01','method':'transferencia','payer':{'type':'owner','id':owner},
            'external_source':'synthetic','external_key':'refund-source','treasury_reference':'erp4-treasury:'+str(self.account)}
        cid=ReceivablesService(self.db).collection_record(self.session,replace(self.e('collection.record',payload),command='erp3.collection.record'))['entity']['id']
        mid=self.imported([self.row('-4000')])[0]
        self.match([{'movement_id':mid,'action':'refund','amount_cents':'-4000','refund_payload':{
            'collection_id':cid,'beneficiary':{'type':'owner','id':owner},'effective_on':'2026-09-01','amount_cents':'4000'}}])
        from erp_core.receivables_projection import collection_balance
        self.assertEqual(collection_balance(self.conn,self.community,cid)['available_cents'],'6000')
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_reintegros').fetchone()[0],1)

    def test_36_proposals_respect_partial_remaining_and_window(self):
        cid=self.collection();mid=self.imported([self.row()])[0]
        self.match([{'movement_id':mid,'action':'link_collection','fact_id':cid,'amount_cents':'4000'}])
        p=self.service.proposals_get(self.session,self.q('proposals.get',{'movement_id':mid,'window_days':0}))['entity']
        self.assertEqual(p['items'][0]['amount_cents'],'6000')
        with self.assertRaises(ContractError):self.service.proposals_get(self.session,self.q('proposals.get',{'movement_id':mid,'window_days':91}))

    def test_37_mapping_profiles_version_and_changed_columns(self):
        raw=b'Fecha;Importe;Concepto\n2026-09-01;100.00;Sintetico\n'
        file={'treasury_id':self.account,'profile':'tabular-v1','filename':'profile.csv','content_base64':base64.b64encode(raw).decode()}
        analyzed=self.service.statement_analyze(self.session,self.e('statement.analyze',file))['entity']
        p={'treasury_id':self.account,'name':'Perfil sintetico','kind':'mapping','config':{'header_signature':analyzed['header_signature'],
            'mapping':{'operation_on':0,'amount':1,'concept':2},'options':{'decimal_separator':'.','date_format':'iso'}}}
        saved=self.service.profile_save(self.session,self.e('profile.save',p,0))['entity']
        found=self.service.statement_analyze(self.session,self.e('statement.analyze',file))['entity']['saved_profiles']
        self.assertEqual(found[0]['id'],saved['id'])
        with self.assertRaises(ConflictError):self.service.profile_save(self.session,self.e('profile.save',p,0))
        self.service.profile_save(self.session,self.e('profile.save',p,1))
        self.assertEqual(len(self.service.statement_analyze(self.session,self.e('statement.analyze',file))['entity']['saved_profiles']),1)
        changed={**file,'content_base64':base64.b64encode(raw.replace(b'Concepto',b'Referencia')).decode()}
        self.assertEqual(self.service.statement_analyze(self.session,self.e('statement.analyze',changed))['entity']['saved_profiles'],[])

    def test_38_matching_profile_is_used_but_never_an_automatic_confirmation(self):
        cid=self.collection();mid=self.imported([self.row(on='2026-09-03')])[0]
        p={'treasury_id':self.account,'name':'Ventana sintetica','kind':'matching','config':{'window_days':0}}
        saved=self.service.profile_save(self.session,self.e('profile.save',p,0))['entity']
        candidates=self.service.proposals_get(self.session,self.q('proposals.get',{'movement_id':mid}))['entity']
        self.assertEqual(candidates['items'],[]);self.assertEqual(candidates['profile_id'],saved['id'])
        self.assertEqual(self.service.proposals_get(self.session,self.q('proposals.get',{'movement_id':mid,'window_days':7}))['entity']['items'][0]['fact_id'],cid)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_conciliacion_componentes').fetchone()[0],0)

    def test_39_pdf_excel_summary_and_pending_are_distinct_from_debt(self):
        mid=self.imported([self.row()],{'start_on':'2026-09-01','end_on':'2026-09-30','opening_cents':'0','closing_cents':'10000','complete':True})[0]
        self.service.pending_assign(self.session,self.e('pending.assign',{'movement_id':mid,'user_id':self.uid,'review_on':'2026-10-01','note':'Identificacion pendiente'}))
        p={'treasury_id':self.account,'start_on':'2026-09-01','end_on':'2026-09-30'}
        pdf=self.service.report_export(self.session,self.e('report.export',{**p,'format':'pdf'}))['entity']
        raw=base64.b64decode(pdf['content_base64']);self.assertTrue(raw.startswith(b'%PDF-'))
        (self.work/'report.pdf').write_bytes(raw)
        xlsx=self.service.report_export(self.session,self.e('report.export',{**p,'format':'xlsx'}))['entity']
        import io
        from openpyxl import load_workbook
        book=load_workbook(io.BytesIO(base64.b64decode(xlsx['content_base64'])))
        self.assertEqual(book.sheetnames,['Conciliacion','Resumen','Pendientes'])
        self.assertEqual(book['Pendientes']['F2'].value,'Identificacion pendiente')
        summary=self.service.report_get(self.session,self.q('report.get',p))['entity']['summary']
        self.assertEqual(summary['unidentified_cents'],'10000');self.assertEqual(summary['unapplied_funds_cents'],'0')
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_cobros').fetchone()[0],0)

    def test_40_cross_profile_alias_adds_evidence_not_money(self):
        first=self.imported([self.row()])[0]
        changed={**self.row(),'namespace':'second-reviewed-profile','concept':'Otra descripcion'}
        p=self.service.statement_preview(self.session,self.e('statement.preview',{'treasury_id':self.account,'profile':'manual-v1','rows':[changed]}))['entity']
        confirmed=self.service.statement_confirm(self.session,self.e('statement.confirm',{'id':p['id'],'decisions':{'1':{
            'action':'alias','id':first,'reason':'Correspondencia de identidad acreditada'}}},1))['entity']
        self.assertEqual(confirmed['movement_ids'],[first])
        self.assertEqual(self.imported([changed]),[first])
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_banco_movimientos').fetchone()[0],1)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_cobros').fetchone()[0],0)

    def test_41_outgoing_proposal_links_existing_return_without_new_debt(self):
        cid=self.collection();mid=self.imported([self.row('-4000')])[0]
        self.match([{'movement_id':mid,'action':'return','amount_cents':'-4000','return_payload':{
            'collection_id':cid,'effective_on':'2026-09-01','free_cents':'4000','reversals':[],'external_key':'return-proposal'}}])
        returned=self.conn.execute('SELECT id FROM erp_devoluciones ORDER BY id DESC LIMIT 1').fetchone()[0]
        match=self.conn.execute('SELECT id FROM erp_conciliaciones ORDER BY id DESC LIMIT 1').fetchone()[0]
        self.service.match_reverse(self.session,self.e('match.reverse',{'id':match}))
        candidate=self.service.proposals_get(self.session,self.q('proposals.get',{'movement_id':mid}))['entity']['items'][0]
        self.assertEqual((candidate['action'],candidate['fact_id'],candidate['amount_cents']),('link_return',returned,'-4000'))
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_devoluciones').fetchone()[0],1)

    def test_42_rotated_keys_preserve_identity_and_restore_ciphertext(self):
        import os,hashlib
        config={'treasury_id':self.account,'name':'Perfil tras rotacion','kind':'matching','config':{'window_days':7}}
        self.service.profile_save(self.session,self.e('profile.save',config,0))
        first=self.imported([self.row()])[0]
        original=json.loads(self.vault.path.read_text())
        original['encryption_keys']['enc-2']=base64.b64encode(os.urandom(32)).decode()
        original['index_keys']['idx-2']=base64.b64encode(os.urandom(32)).decode()
        original.update(active_encryption_key='enc-2',active_index_key='idx-2')
        self.vault.path.write_text(json.dumps(original));self.vault=BankVault(self.vault.path)
        from erp_core.database import write_transaction
        with write_transaction(self.conn):self.vault.rewrap(self.conn,self.community)
        self.service=ReconciliationService(self.db,vault=self.vault)
        self.assertEqual(self.service.profile_save(self.session,self.e('profile.save',config,1))['entity']['version'],2)
        self.assertEqual(self.imported([{**self.row(),'concept':'Nueva evidencia tras rotacion'}]),[first])
        target=self.work/'restored.db';key=self.work/'restored-custody.key'
        fd=os.open(key,os.O_WRONLY|os.O_CREAT|os.O_EXCL,0o600)
        with os.fdopen(fd,'wb') as stream:stream.write(self.vault.path.read_bytes())
        with closing(sqlite3.connect(target)) as destination:self.conn.backup(destination)
        restored=ReconciliationService(target,vault=BankVault(key))
        self.assertEqual(restored.movement_get(self.session,self.q('movement.get',{'id':first}))['entity']['amount_cents'],'10000')
        with closing(connect(target)) as conn:
            self.assertEqual(conn.execute('PRAGMA integrity_check').fetchone()[0],'ok');self.assertEqual(conn.execute('PRAGMA foreign_key_check').fetchall(),[])
            self.assertEqual(conn.execute('SELECT count(*) FROM erp_banco_movimientos').fetchone()[0],1)
        with target.open('rb') as stream:digest=hashlib.file_digest(stream,'sha256').hexdigest()
        (self.work/'restore-proof.json').write_text(json.dumps({'ok':True,'sha256':digest,'keys_restored':True}))

    def test_43_link_events_do_not_duplicate_monetary_effect(self):
        cid=self.collection();mid=self.imported([self.row()])[0]
        self.match([{'movement_id':mid,'action':'link_collection','fact_id':cid,'amount_cents':'10000'}])
        event=self.conn.execute("SELECT event_id,payload_json FROM erp_outbox WHERE event_type='erp5.reconciliation.confirmed' ORDER BY id_outbox DESC LIMIT 1").fetchone()
        payload=json.loads(event['payload_json']);link=payload['links'][0]
        self.assertEqual(payload['effect'],'link');self.assertEqual(link['economic_fact_id'],'erp3:collection:'+str(cid))
        self.assertIsNone(link['economic_event_id']) # ERP 3 collection rows have no direct event FK; do not infer one.
        inbox=set();effects=set()
        for _ in range(2):
            key=(self.community,event['event_id'],link['component_id']);inbox.add(key)
            if payload['effect']=='monetary':effects.add(link['economic_fact_id'])
        self.assertEqual(len(inbox),1);self.assertEqual(len(effects),0)
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_cobros').fetchone()[0],1)

    def test_44_legacy_rows_require_community_and_exact_amount(self):
        with self.conn:
            known=self.conn.execute('INSERT INTO cf_extractos_banco_lineas(line_hash,importe,fecha,id_comunidad) VALUES (?,?,?,?)',
                ('synthetic-scoped',100,'2026-09-01',self.community)).lastrowid
            unknown=self.conn.execute('INSERT INTO cf_extractos_banco_lineas(line_hash,importe,fecha,id_comunidad) VALUES (?,?,?,NULL)',
                ('synthetic-unattributed',100,'2026-09-01')).lastrowid
        row={'table':'cf_extractos_banco_lineas','id':known,'movement':self.row()}
        proposal=self.service.legacy_preview(self.session,self.e('legacy.preview',{'treasury_id':self.account,'rows':[row]}))['entity']
        self.assertEqual(proposal['rows'][0]['quality'],'observada')
        from erp_core.errors import NotFoundError
        with self.assertRaises(NotFoundError):self.service.legacy_preview(self.session,self.e('legacy.preview',{'treasury_id':self.account,'rows':[{**row,'id':unknown}]}))
        with self.assertRaises(ConflictError):self.service.legacy_preview(self.session,self.e('legacy.preview',{'treasury_id':self.account,'rows':[{**row,'movement':self.row('9999')}]}))
        self.assertEqual(self.conn.execute('SELECT count(*) FROM erp_cobros').fetchone()[0],0)


if __name__=='__main__':
    print('workspace='+str(WORK),flush=True)
    unittest.main(verbosity=2)
