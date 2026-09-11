"""Install pinned pure-Python export dependencies inside this application only."""
import hashlib
import io
import json
from pathlib import Path
import sys
import urllib.request
import zipfile

ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / 'server' / '_python_packages'
WHEELS = (
    ('openpyxl', 'https://files.pythonhosted.org/packages/c0/da/977ded879c29cbd04de313843e76868e6e13408a94ed6b987245dc7c8506/openpyxl-3.1.5-py2.py3-none-any.whl', '5282c12b107bffeef825f4617dc029afaf41d0ea60823bbb665ef3079dc79de2'),
    ('et_xmlfile', 'https://files.pythonhosted.org/packages/c1/8b/5fe2cc11fee489817272089c4203e679c63b570a5aaeb18d852ae3cbba6a/et_xmlfile-2.0.0-py3-none-any.whl', '7a91720bc756843502c3b7504c77b8fe44217c85c537d85037f0f536151b2caa'),
)
for name, url, expected in WHEELS:
    with urllib.request.urlopen(url, timeout=30) as response:
        content = response.read(2 * 1024 * 1024)
    assert hashlib.sha256(content).hexdigest() == expected, name + ': checksum mismatch'
    with zipfile.ZipFile(io.BytesIO(content)) as bundle:
        for entry in bundle.infolist():
            destination = (TARGET / entry.filename).resolve()
            assert destination.is_relative_to(TARGET.resolve()), 'Invalid wheel path'
        bundle.extractall(TARGET)
sys.path.insert(0, str(TARGET))
from openpyxl import Workbook, load_workbook
stream = io.BytesIO()
book = Workbook()
book.active.append(['1234567890123.45'])
book.save(stream)
stream.seek(0)
assert load_workbook(stream).active['A1'].value == '1234567890123.45'
print(json.dumps({'dependencies': str(TARGET), 'xlsx_exact_text': True}))
