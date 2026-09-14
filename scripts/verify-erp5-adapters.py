"""Synthetic, bounded statement fixtures; never access live bank data."""
import base64
import io
from pathlib import Path
import sys
import unittest
import zipfile

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'server'))
from erp_core.reconciliation_adapters import parse,amount
from erp_core.errors import ContractError


def file_payload(raw,name,profile='tabular-v1',**extra):
    return {'content_base64':base64.b64encode(raw).decode(),'filename':name,'profile':profile,**extra}


class AdapterTests(unittest.TestCase):
    def csv(self,content='Fecha;Valor;Importe;Concepto\n01/09/2026;02/09/2026;1.234,56;Prueba\n'):
        return file_payload(content.encode(),'synthetic.csv',mapping={'operation_on':0,'value_on':1,'amount':2,'concept':3},
                            options={'decimal_separator':',','date_format':'dmy'})

    def test_01_csv_mapping_exact(self):
        result=parse(self.csv());self.assertEqual(result['rows'][0]['amount_cents'],'123456');self.assertEqual(result['errors'],[])

    def test_02_csv_row_errors_not_silent(self):
        result=parse(self.csv('Fecha;Valor;Importe;Concepto\n01/09/2026;;0,001;Prueba\n'))
        self.assertEqual(len(result['errors']),1);self.assertEqual(result['rows'],[])

    def test_03_analyze_without_mapping(self):
        p=self.csv();p.pop('mapping');self.assertEqual(parse(p)['headers'][0],'Fecha')

    def test_04_credit_debit(self):
        p=file_payload(b'Date;Debit;Credit\n2026-09-01;10.20;0\n2026-09-02;0;30.50\n','synthetic.csv',
                       mapping={'operation_on':0,'debit':1,'credit':2})
        self.assertEqual([r['amount_cents'] for r in parse(p)['rows']],['-1020','3050'])

    def test_05_xlsx_formula_cannot_be_money(self):
        from openpyxl import Workbook
        book=Workbook();sheet=book.active;sheet.append(['Fecha','Importe']);sheet.append(['2026-09-01','=1+1'])
        stream=io.BytesIO();book.save(stream)
        result=parse(file_payload(stream.getvalue(),'synthetic.xlsx',mapping={'operation_on':0,'amount':1}))
        self.assertEqual(len(result['errors']),1)

    def test_06_external_links_rejected(self):
        stream=io.BytesIO()
        with zipfile.ZipFile(stream,'w') as archive:archive.writestr('xl/externalLinks/externalLink1.xml','<root/>')
        with self.assertRaises(ContractError):parse(file_payload(stream.getvalue(),'hostile.xlsx'))

    def test_07_macros_rejected(self):
        stream=io.BytesIO()
        with zipfile.ZipFile(stream,'w') as archive:archive.writestr('xl/vbaProject.bin',b'fixture')
        with self.assertRaises(ContractError):parse(file_payload(stream.getvalue(),'hostile.xlsx'))

    def xml(self,profile='camt.053.001.08',detail='100.00',status='BOOK'):
        block='Stmt' if profile.startswith('camt.053') else 'Ntfctn';parent='BkToCstmrStmt' if block=='Stmt' else 'BkToCstmrDbtCdtNtfctn'
        balances='''<Bal><Tp><CdOrPrtry><Cd>OPBD</Cd></CdOrPrtry></Tp><Amt Ccy="EUR">0.00</Amt><CdtDbtInd>CRDT</CdtDbtInd><Dt><Dt>2026-08-31</Dt></Dt></Bal>
            <Bal><Tp><CdOrPrtry><Cd>CLBD</Cd></CdOrPrtry></Tp><Amt Ccy="EUR">100.00</Amt><CdtDbtInd>CRDT</CdtDbtInd><Dt><Dt>2026-09-01</Dt></Dt></Bal>''' if block=='Stmt' else ''
        return file_payload(f'''<Document xmlns="urn:iso:std:iso:20022:tech:xsd:{profile}"><{parent}><GrpHdr><MsgId>SYNTHETIC-1</MsgId><CreDtTm>2026-09-02T10:00:00Z</CreDtTm></GrpHdr><{block}>
            <Id>SYNTHETIC-BLOCK</Id><CreDtTm>2026-09-02T10:00:00Z</CreDtTm>
            <Acct><Id><IBAN>ES9121000418450200051332</IBAN></Id></Acct>
            {balances}
            <Ntry><Amt Ccy="EUR">100.00</Amt><CdtDbtInd>CRDT</CdtDbtInd><Sts><Cd>{status}</Cd></Sts>
            <BookgDt><Dt>2026-09-01</Dt></BookgDt><ValDt><Dt>2026-09-02</Dt></ValDt><AcctSvcrRef>SYNTHETIC</AcctSvcrRef>
            <BkTxCd><Prtry><Cd>SYNTHETIC</Cd></Prtry></BkTxCd>
            <NtryDtls><TxDtls><AmtDtls><TxAmt><Amt Ccy="EUR">{detail}</Amt></TxAmt></AmtDtls><RmtInf><Ustrd>Prueba</Ustrd></RmtInf></TxDtls></NtryDtls>
            </Ntry></{block}></{parent}></Document>'''.encode(),'synthetic.xml',profile)

    def test_08_camt_aggregate_not_double_sum(self):
        r=parse(self.xml());self.assertEqual(len(r['rows']),1);self.assertEqual(r['rows'][0]['amount_cents'],'10000')

    def test_09_camt_details_must_balance(self):
        with self.assertRaises(ContractError):parse(self.xml(detail='99.99'))

    def test_10_camt054_pending(self):
        self.assertEqual(parse(self.xml('camt.054.001.08',status='PDNG'))['rows'][0]['source_state'],'pending')

    def test_11_xml_dtd_rejected(self):
        with self.assertRaises(ContractError):parse(file_payload(b'<!DOCTYPE a [<!ENTITY x SYSTEM "file:///etc/passwd">]><a>&x;</a>','hostile.xml','camt.053.001.08'))

    def norm(self):
        ccc='210004180200051332'
        head='11'+ccc+'260901'+'260930'+'2'+'00000000001000'+'978'+' '*30
        entry='22'+' '*8+'260901'+'260902'+'00000'+'2'+'00000000010000'+' '*38
        tail='33'+ccc+'00000'+'00000000000000'+'00001'+'00000000010000'+'2'+'00000000011000'+'978'+' '*4
        assert len(head)==len(entry)==len(tail)==80
        return file_payload(('\n'.join([head,entry,tail])+'\n').encode(),'synthetic.txt','cuaderno43-v1')

    def test_12_norm43_balances(self):
        result=parse(self.norm());self.assertEqual(result['coverage']['closing_cents'],'11000');self.assertEqual(result['account_ccc'],'210004180200051332')

    def test_13_norm43_wrong_totals(self):
        p=self.norm();raw=base64.b64decode(p['content_base64']).replace(b'00000000011000',b'00000000011001')
        with self.assertRaises(ContractError):parse(file_payload(raw,'synthetic.txt','cuaderno43-v1'))

    def test_14_unknown_profile(self):
        with self.assertRaises(ContractError):parse({'profile':'unknown','rows':[]})

    def test_15_excel_numeric_cells_ignore_text_decimal_separator(self):
        from openpyxl import Workbook
        book=Workbook();sheet=book.active;sheet.append(['Fecha','Importe']);sheet.append(['2026-09-01',100]);sheet.append(['2026-09-02','123,45'])
        stream=io.BytesIO();book.save(stream)
        result=parse(file_payload(stream.getvalue(),'synthetic.xlsx',mapping={'operation_on':0,'amount':1},options={'decimal_separator':','}))
        self.assertEqual([r['amount_cents'] for r in result['rows']],['10000','12345'])

    def test_16_xls_numeric_mapping(self):
        import xlwt
        book=xlwt.Workbook();sheet=book.add_sheet('Datos');sheet.write(0,0,'Fecha');sheet.write(0,1,'Importe');sheet.write(1,0,'2026-09-01');sheet.write(1,1,100)
        stream=io.BytesIO();book.save(stream)
        result=parse(file_payload(stream.getvalue(),'synthetic.xls',mapping={'operation_on':0,'amount':1},options={'decimal_separator':','}))
        self.assertEqual(result['rows'][0]['amount_cents'],'10000')

    def test_17_xls_formulas_rejected(self):
        import xlwt
        book=xlwt.Workbook();sheet=book.add_sheet('Datos');sheet.write(0,0,'Fecha');sheet.write(0,1,'Importe');sheet.write(1,0,'2026-09-01');sheet.write(1,1,xlwt.Formula('1+1'))
        stream=io.BytesIO();book.save(stream)
        with self.assertRaises(ContractError):parse(file_payload(stream.getvalue(),'synthetic.xls',mapping={'operation_on':0,'amount':1}))

    def test_18_camt_direct_amount_details_are_checked(self):
        p=self.xml();raw=base64.b64decode(p['content_base64']).replace(b'<AmtDtls><TxAmt><Amt Ccy="EUR">100.00</Amt></TxAmt></AmtDtls>',b'<Amt Ccy="EUR">100.00</Amt>')
        self.assertEqual(parse(file_payload(raw,'direct.xml','camt.053.001.08'))['rows'][0]['amount_cents'],'10000')
        with self.assertRaises(ContractError):parse(file_payload(raw.replace(b'<TxDtls><Amt Ccy="EUR">100.00',b'<TxDtls><Amt Ccy="EUR">99.99'),'direct.xml','camt.053.001.08'))

    def test_19_camt_signed_net_details_not_double_counted(self):
        p=self.xml();raw=base64.b64decode(p['content_base64'])
        start=raw.index(b'<NtryDtls>');end=raw.index(b'</NtryDtls>')+len(b'</NtryDtls>')
        details=b'<NtryDtls><TxDtls><Amt Ccy="EUR">120.00</Amt><CdtDbtInd>CRDT</CdtDbtInd></TxDtls><TxDtls><Amt Ccy="EUR">20.00</Amt><CdtDbtInd>DBIT</CdtDbtInd></TxDtls></NtryDtls>'
        result=parse(file_payload(raw[:start]+details+raw[end:],'net.xml','camt.053.001.08'))
        self.assertEqual(len(result['rows']),1);self.assertEqual(result['rows'][0]['amount_cents'],'10000')

    def test_20_camt_intra_day_balances_do_not_certify_whole_days(self):
        raw=base64.b64decode(self.xml()['content_base64'])
        def interval(beginning,ending):return raw.replace(b'<Acct>',f'<FrToDt><FrDtTm>{beginning}</FrDtTm><ToDtTm>{ending}</ToDtTm></FrToDt><Acct>'.encode(),1)
        partial=parse(file_payload(interval('2026-09-01T10:00:00Z','2026-09-01T18:00:00Z'),'partial.xml','camt.053.001.08'))
        self.assertFalse(partial['coverage']['complete'])
        complete=parse(file_payload(interval('2026-09-01T00:00:00Z','2026-09-02T00:00:00Z'),'complete.xml','camt.053.001.08'))
        self.assertTrue(complete['coverage']['complete']);self.assertEqual(complete['coverage']['end_on'],'2026-09-01')


    def test_21_multiblock_camt_explicit_selection_preserves_original_hash(self):
        import copy
        from lxml import etree
        p=self.xml();root=etree.fromstring(base64.b64decode(p['content_base64']));parent=root[0]
        original=parent[1];second=copy.deepcopy(original);parent.append(second)
        raw=etree.tostring(root);p=file_payload(raw,'multi.xml','camt.053.001.08')
        manifest=parse(p);self.assertTrue(manifest['requires_block']);self.assertEqual(manifest['rows'],[])
        self.assertNotIn('ES912100',str(manifest['blocks']))
        first=parse({**p,'options':{'block_index':0}});second=parse({**p,'options':{'block_index':1}})
        self.assertEqual(first['file_hash'],second['file_hash']);self.assertNotEqual(first['source_hash'],second['source_hash'])
        self.assertEqual(first['rows'][0]['amount_cents'],'10000');self.assertEqual(len(second['rows']),1)
        for index in (-1,2,'0',True):
            with self.assertRaises(ContractError):parse({**p,'options':{'block_index':index}})

    def test_22_multiblock_norm43_requires_account_selection(self):
        raw=base64.b64decode(self.norm()['content_base64']);p=file_payload(raw+raw,'multi.txt','cuaderno43-v1')
        self.assertTrue(parse(p)['requires_block'])
        first=parse({**p,'options':{'block_index':0}});second=parse({**p,'options':{'block_index':1}})
        self.assertNotEqual(first['source_hash'],second['source_hash']);self.assertEqual(second['coverage']['closing_cents'],'11000')
        with self.assertRaises(ContractError):parse(file_payload(raw.splitlines()[0]+b'\n'+raw,'bad.txt','cuaderno43-v1'))


if __name__=='__main__':unittest.main(verbosity=2)
